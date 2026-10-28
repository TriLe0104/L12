"""Argus test catalog, runs, logs, and reports — grouped by the component they exercise."""

from __future__ import annotations

import json
from typing import Any

from . import cluster_controller

COMPONENTS: list[dict[str, str]] = [
    {"id": "gpu", "label": "GPU", "blurb": "Burn, DCGM, SMI, GEMM"},
    {"id": "interconnect", "label": "Interconnect", "blurb": "NCCL, NVLink, C2C"},
    {"id": "network", "label": "Network", "blurb": "Ping, IB, connectivity"},
    {"id": "pcie", "label": "PCIe", "blurb": "lspci and link checks"},
    {"id": "memory", "label": "Memory", "blurb": "HBM, DRAM, GPU mem"},
    {"id": "cpu", "label": "CPU", "blurb": "Stress-ng, Tegra CPU"},
    {"id": "storage", "label": "Storage", "blurb": "SSD / NVMe"},
    {"id": "thermal", "label": "Thermal", "blurb": "Heat and cooling"},
    {"id": "system", "label": "System", "blurb": "L10 / L11 tray diagnostics"},
]


def classify(test: dict[str, Any]) -> str:
    cid = str(test.get("id") or "").lower()
    clf = str(test.get("test_classifier") or "").lower()
    name = str(test.get("name") or "").lower()
    tags = " ".join(str(t) for t in (test.get("tags") or [])).lower()
    blob = f"{cid} {clf} {name} {tags}"
    if any(x in blob for x in ("ssd", "nvme", "e1s")):
        return "storage"
    if any(x in blob for x in ("network", "ping", "ib_", "rdma", "perftest", "ib_write", "ib_read", "ib_send")):
        return "network"
    if any(x in blob for x in ("nccl", "nvl", "nvlink", "c2c", "clink", "nvl-mapper")):
        return "interconnect"
    if any(x in blob for x in ("pcie", "lspci")):
        return "pcie"
    if any(x in blob for x in ("hbm", "memory", "memsweep", "gpumem", "tegra_memory", "cpu_memory")):
        return "memory"
    if any(x in blob for x in ("stress-ng", "tegra_cpu", "cpu_gpu")):
        return "cpu"
    if any(x in blob for x in ("thermal",)):
        return "thermal"
    if any(x in blob for x in ("dcgm", "gpu-burn", "gpu burn", "nvidia-smi", "gpustress", "gemm")):
        return "gpu"
    if any(x in blob for x in ("l10", "l11", "diag")):
        return "system"
    return "system"


def _public_test(raw: dict[str, Any]) -> dict[str, Any]:
    subs = raw.get("l10_subtests") if isinstance(raw.get("l10_subtests"), list) else []
    return {
        "id": raw.get("id"),
        "name": raw.get("id"),
        "label": raw.get("name") or raw.get("id"),
        "detail": raw.get("description") or "",
        "classifier": raw.get("test_classifier"),
        "component": classify(raw),
        "timeout_seconds": raw.get("timeout_seconds"),
        "parent_test_id": raw.get("parent_test_id"),
        "subtests": [
            {"id": s.get("id"), "label": s.get("name") or s.get("id")}
            for s in subs
            if isinstance(s, dict)
        ],
        "source": "argus",
        "kind": classify(raw),
        "gpu_request": 0,
        "gpu_mem_gb_request": 0,
        "pods_requested": 1,
    }


def catalog() -> dict[str, Any]:
    if not cluster_controller.enabled():
        return {"source": "none", "components": COMPONENTS, "tests": [], "error": "Argus URL is not set"}
    try:
        raw = cluster_controller.tests_list()
    except Exception as exc:
        return {"source": "argus", "components": COMPONENTS, "tests": [], "error": str(exc)[:240]}
    tests = [_public_test(t) for t in raw if not t.get("parent_test_id")]
    by_comp: dict[str, int] = {}
    for t in tests:
        by_comp[str(t["component"])] = by_comp.get(str(t["component"]), 0) + 1
    return {
        "source": "argus",
        "components": [{**c, "count": by_comp.get(c["id"], 0)} for c in COMPONENTS],
        "tests": tests,
        "error": None,
    }


def active_runs() -> dict[str, Any]:
    if not cluster_controller.enabled():
        return {"active_runs": []}
    try:
        return cluster_controller.active_test_runs()
    except Exception as exc:
        return {"active_runs": [], "error": str(exc)[:240]}


def _report_name(path: str) -> str:
    return (path or "").rstrip("/").split("/")[-1] or "report"


def _public_report(raw: dict[str, Any]) -> dict[str, Any]:
    path = str(raw.get("path") or "")
    kind = str(raw.get("type") or "").strip().lower()
    if not kind:
        low = path.lower()
        if low.endswith(".pdf"):
            kind = "pdf"
        elif low.endswith((".xlsx", ".xls")):
            kind = "xlsx"
        elif low.endswith(".json"):
            kind = "json"
    runs = raw.get("test_runs")
    if not isinstance(runs, list):
        runs = []
    created = raw.get("date_created")
    if created is not None and not isinstance(created, str):
        created = str(created)
    return {
        "serial_number": raw.get("serial_number"),
        "type": kind,
        "path": path,
        "name": _report_name(path),
        "date_created": created,
        "test_runs": [r for r in runs if r is not None],
    }


def reports(serial: str | None = None, name: str | None = None) -> list[dict[str, Any]]:
    if not cluster_controller.enabled():
        return []
    rows = cluster_controller.reports(serial, name)
    return [_public_report(r) for r in rows if isinstance(r, dict)]


def list_reports(serial: str | None = None, name: str | None = None) -> dict[str, Any]:
    if not cluster_controller.enabled():
        return {"reports": [], "count": 0, "error": "Argus URL is not set"}
    try:
        rows = reports(serial, name)
        return {"reports": rows, "count": len(rows), "error": None}
    except Exception as exc:
        return {"reports": [], "count": 0, "error": str(exc)[:240]}


def rack_results(serial: str, name: str) -> Any:
    return cluster_controller.rack_test_results(serial, name)


def latest_report(serial: str, name: str) -> Any:
    return cluster_controller.rack_latest_report(serial, name)


def node_logs(bmc_mac: str) -> Any:
    return cluster_controller.node_logs(bmc_mac)


def logs_index() -> Any:
    return cluster_controller.logs_all()


def read_log(path: str) -> Any:
    return cluster_controller.read_log_file(path)


def _guess_mime(path: str, data: bytes) -> str:
    low = (path or "").lower()
    if data.startswith(b"%PDF") or low.endswith(".pdf"):
        return "application/pdf"
    if data[:2] == b"PK" or low.endswith(".xlsx"):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if low.endswith(".json"):
        return "application/json"
    if low.endswith((".txt", ".log", ".csv")):
        return "text/plain"
    return "application/octet-stream"


_PREVIEW_SHEETS = 8
_PREVIEW_ROWS = 120
_PREVIEW_COLS = 40
_PREVIEW_JSON = 200_000


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def preview_bytes(data: bytes, name: str, mime: str) -> dict[str, Any]:
    low = (name or "").lower()
    kind_mime = (mime or "").split(";")[0].strip().lower()
    if data.startswith(b"%PDF") or low.endswith(".pdf"):
        return {"kind": "pdf", "filename": name, "mime": "application/pdf"}
    if data[:2] == b"PK" or low.endswith((".xlsx", ".xls")) or "spreadsheet" in kind_mime:
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        try:
            sheets: list[dict[str, Any]] = []
            for sheet in wb.worksheets[:_PREVIEW_SHEETS]:
                rows: list[list[str]] = []
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    if i >= _PREVIEW_ROWS:
                        break
                    cells = [_cell(c) for c in (row or ())[:_PREVIEW_COLS]]
                    rows.append(cells)
                sheets.append({"name": sheet.title, "rows": rows})
            truncated = len(wb.worksheets) > _PREVIEW_SHEETS
        finally:
            wb.close()
        return {
            "kind": "xlsx",
            "filename": name,
            "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "sheets": sheets,
            "truncated": truncated,
        }
    text = data.decode("utf-8", "replace")
    if low.endswith(".json") or kind_mime in ("application/json", "text/json") or text[:1] in "{[":
        try:
            payload = json.loads(text)
            pretty = json.dumps(payload, indent=2, default=str)
        except Exception:
            pretty = text
        truncated = len(pretty) > _PREVIEW_JSON
        if truncated:
            pretty = pretty[:_PREVIEW_JSON] + "\n… [truncated]"
        return {"kind": "json", "filename": name, "mime": "application/json", "text": pretty, "truncated": truncated}
    truncated = len(text) > _PREVIEW_JSON
    if truncated:
        text = text[:_PREVIEW_JSON] + "\n… [truncated]"
    return {"kind": "text", "filename": name, "mime": kind_mime or "text/plain", "text": text, "truncated": truncated}


def preview_file(path: str) -> dict[str, Any]:
    data, mime, name = read_file(path)
    return preview_bytes(data, name, mime)


def read_file(path: str) -> tuple[bytes, str, str]:
    raw, ctype = cluster_controller.read_file_bytes(path)
    name = path.rstrip("/").split("/")[-1] or "report"
    if raw.startswith(b"%PDF"):
        return raw, "application/pdf", name if name.lower().endswith(".pdf") else f"{name}.pdf"
    mime = (ctype or "").split(";")[0].strip().lower()
    if mime in ("application/json", "text/json") and raw[:1] in (b"{", b"["):
        try:
            payload = json.loads(raw.decode("utf-8", "replace"))
        except Exception:
            payload = None
        if isinstance(payload, dict):
            for key in ("content", "data", "file", "body", "bytes"):
                blob = payload.get(key)
                if isinstance(blob, str) and blob.startswith("%PDF"):
                    data = blob.encode("utf-8")
                    return data, "application/pdf", name
                if isinstance(blob, str) and len(blob) > 20:
                    import base64

                    try:
                        data = base64.b64decode(blob, validate=False)
                    except Exception:
                        data = b""
                    if data.startswith(b"%PDF"):
                        return data, "application/pdf", name
                    if data[:2] == b"PK":
                        return data, _guess_mime(name, data), name
            text = json.dumps(payload, indent=2, default=str).encode("utf-8")
            return text, "application/json", name if name.lower().endswith(".json") else f"{name}.json"
    if mime in ("", "application/octet-stream"):
        mime = _guess_mime(path, raw)
    return raw, mime, name


def launch(tests: list[str], racks: list[dict[str, Any]], l10_subtests: dict[str, list[str]] | None = None) -> Any:
    return cluster_controller.run_test(tests, racks, l10_subtests)
