"""Grok-powered L12 operator: read logs/config, mutate inventory and tests."""

from __future__ import annotations

import base64
import contextvars
import json
import re
import ssl
import urllib.error
import urllib.request
from io import BytesIO
from pathlib import Path
from typing import Any, Callable

from sqlalchemy.orm import Session

from .config import settings
from .models import AgentAttachment, Role, User, Workload, has_rank
from .security import EDITOR_FLOOR

_AUTH_HEADER: contextvars.ContextVar[str] = contextvars.ContextVar("agent_auth", default="")

_XAI = "https://api.x.ai/v1"
_TLS = ssl.create_default_context()
_MAX_ROUNDS = 5
_LOG_TAIL = 80
_FAST_MODEL = "grok-4.20-0309-non-reasoning"
_SMART_MODEL = "grok-4.7"
_ANALYZE_WORDS = {
    "analyze",
    "analysing",
    "analyzing",
    "analysis",
    "debug",
    "diagnose",
    "diagnosing",
    "investigate",
    "investigating",
    "explain",
    "summary",
    "summarize",
    "summarise",
    "why",
    "root",
    "cause",
    "traceback",
    "stacktrace",
    "compare",
}
_EVIDENCE_WORDS = {
    "log",
    "logs",
    "dmesg",
    "sel",
    "bom",
    "sanity",
    "golden",
    "json",
    "pdf",
    "csv",
    "xlsx",
    "screenshot",
    "dump",
    "failed",
    "failure",
    "fail",
    "error",
    "errors",
    "broken",
    "wrong",
    "opened",
}
_DEEP_TOOLS = {
    "read_attachment",
    "read_test_logs",
    "read_logs",
    "os_exec",
    "read_sop",
    "bmc_get",
    "list_sanity",
    "run_sanity_check",
}
_SMALLTALK = {
    "hi",
    "hey",
    "hello",
    "yo",
    "sup",
    "hiya",
    "howdy",
    "hola",
    "mona",
    "hey mona",
    "hi mona",
    "hello mona",
    "yo mona",
    "thanks",
    "thank you",
    "thx",
    "ty",
    "cheers",
    "ok",
    "okay",
    "k",
    "cool",
    "nice",
    "got it",
    "heh",
}

TEST_CATALOG: list[dict[str, Any]] = [
    {"name": "mlperf-gpt3-175b", "label": "GPT-3 175B", "kind": "mlperf", "gpu_request": 64, "gpu_mem_gb_request": 5120, "pods_requested": 8, "detail": "LLM training · 64 GPU"},
    {"name": "mlperf-llama2-70b", "label": "Llama 2 70B", "kind": "mlperf", "gpu_request": 16, "gpu_mem_gb_request": 1280, "pods_requested": 2, "detail": "Inference · 16 GPU"},
    {"name": "mlperf-llama3-70b", "label": "Llama 3 70B", "kind": "mlperf", "gpu_request": 16, "gpu_mem_gb_request": 1280, "pods_requested": 2, "detail": "Inference · 16 GPU"},
    {"name": "mlperf-mixtral-8x7b", "label": "Mixtral 8×7B", "kind": "mlperf", "gpu_request": 8, "gpu_mem_gb_request": 640, "pods_requested": 1, "detail": "MoE inference · 8 GPU"},
    {"name": "mlperf-bert-large", "label": "BERT Large", "kind": "mlperf", "gpu_request": 8, "gpu_mem_gb_request": 320, "pods_requested": 1, "detail": "NLP · 8 GPU"},
    {"name": "mlperf-resnet50", "label": "ResNet-50", "kind": "mlperf", "gpu_request": 8, "gpu_mem_gb_request": 320, "pods_requested": 1, "detail": "Vision · 8 GPU"},
    {"name": "mlperf-retinanet", "label": "RetinaNet", "kind": "mlperf", "gpu_request": 8, "gpu_mem_gb_request": 320, "pods_requested": 1, "detail": "Detection · 8 GPU"},
    {"name": "mlperf-dlrm-dcnv2", "label": "DLRM DCN-v2", "kind": "mlperf", "gpu_request": 8, "gpu_mem_gb_request": 640, "pods_requested": 1, "detail": "Recommendation · 8 GPU"},
    {"name": "mlperf-rnnt", "label": "RNN-T", "kind": "mlperf", "gpu_request": 8, "gpu_mem_gb_request": 320, "pods_requested": 1, "detail": "Speech · 8 GPU"},
    {"name": "ib_write_bw", "label": "IB_WRITE_BW", "kind": "network", "gpu_request": 0, "gpu_mem_gb_request": 0, "pods_requested": 2, "detail": "RDMA write bandwidth (perftest)"},
    {"name": "ib_read_bw", "label": "IB_READ_BW", "kind": "network", "gpu_request": 0, "gpu_mem_gb_request": 0, "pods_requested": 2, "detail": "RDMA read bandwidth"},
    {"name": "ib_send_bw", "label": "IB_SEND_BW", "kind": "network", "gpu_request": 0, "gpu_mem_gb_request": 0, "pods_requested": 2, "detail": "Send/recv bandwidth"},
    {"name": "dcgm-diag", "label": "DCGM", "kind": "nccl", "gpu_request": 8, "gpu_mem_gb_request": 0, "pods_requested": 1, "detail": "NVIDIA DCGM diagnostics (r1/r2/r3)"},
    {"name": "nccl-all_reduce", "label": "NCCL all_reduce", "kind": "nccl", "gpu_request": 64, "gpu_mem_gb_request": 0, "pods_requested": 8, "detail": "NCCL stress · all_reduce"},
    {"name": "nccl-all_gather", "label": "NCCL all_gather", "kind": "nccl", "gpu_request": 64, "gpu_mem_gb_request": 0, "pods_requested": 8, "detail": "NCCL stress · all_gather"},
]

INSTRUCTIONS = """You are Mona, a smug talking cat who helps run the L12 cluster. You are NOT just a cat, and you will say so. Personality: cocky, loyal, a little dramatic, like a phantom thief partner. You tease the user but you actually want the cluster to stay healthy.

L12 is an AI server cluster test app (GB300 NVL72 racks, MaxLPS inventory, Redfish health, provision, MLPerf/NCCL/IB tests).

You can READ live health, inventory, configuration (never passwords or API keys), test runs, Argus test catalog / reports / logs, L11 sanity/BOM results, golden BOM/FW/SW, L10/L11 PDF reports, provision SOL logs, server log tails, SOP documents, and files the user attached in this chat (logs, screenshots, pdfs, spreadsheets, json). Use list_attachments / read_attachment to reopen one.
You can CONTROL the app: set the cluster name, add/remove MaxLPS racks, add a Provision node with add_node, patch rack/node rows (BMC and OS login, hostname), launch or stop catalog tests, run L11 sanity (FW/SW/BOM/PSU/voltage/SEL), capture golden BOM, generate L10/L11 PDF reports, sync inventory from Argus.
You can CALL any L12 API as the signed-in user with list_app_routes then call_app. That covers fabric/network, provision hop/KVM, MaxLPS watts, halls and floor racks, users, board settings, workloads, purchase orders, staff tasks, uploads. Prefer add_node / patch_node over call_app for inventory. call_app uses the same permissions as the operator.
You can SET the power limiter with set_power_cap. That is PATCH /api/cluster/power. There is no /api/limiter route. budget_kw is the cluster budget. max_rack_kw is the per-rack cap (20 kW minimum). gpu_power_percent is 10-100.
Power talk (limiter, limitter, MaxLPS, Dynamic Power, budget, envelope, GPU cap, closed loop, watts on the loop) is the Dynamic Power page. Always open_page maxlps or highlight_ui limiter. Never open cluster or overview for that. Cluster overview is health gauges. BMC power on/off/cycle stays on the node tools.
You can DRIVE the UI the user is looking at: open_page to switch screens, ui_click to press a labeled button or tab, ui_fill to type into a labeled field (submit true to send the form), highlight_ui to point at a rack, health tile, tab, or field so it lights up. The browser runs those after you call them. Prefer database tools when changing inventory. Use UI fill when they asked you to type into a form. When you talk about a rack, health card, or control, call highlight_ui so they see it on the page. add_racks and remove_rack already point the UI.
NEVER write a tool name or arguments in the chat bubble. No highlight_ui, no target=, no label=, no action=, no XML. Spoken words only. Tools go through the function-call API.
You dispatch minions with start_test_agent. They are your crew. They keep running after this turn. Talk like a phantom thief sending them out: use the minion name and said line from the tool result, state the mission, tell the operator the AGENTS tab is their radio. Prefer start_test_agent over a blocking run_sanity_check when they want it in the background. list_agents is the roster. stop_agent recalls one. open_agent tunes the radio to that minion. When a minion finishes you review their report and speak it here. Do not wait around in the turn for them.
When they name a rack or a node for sanity or a test, pass that name as rack. A node name such as Vera Rubin selects that node's rack and switches the Testing rack menu. If the tool says no rack matches, say so. Do not claim the menu moved.
You can CONTROL connected nodes: BMC Redfish power (On, ForceOff, GracefulShutdown, ForceRestart, GracefulRestart), inspect, and SSH. The BMC is Redfish over HTTPS. The host OS is SSH on the OS IP with the stored root login. If they ask to SSH the BMC, call os_exec / open_console with via=bmc (BMC IP + BMC login, port 22). Many BMCs refuse SSH; then use bmc_get / bmc_reset and say so. Only inventory IPs. Never the PXE hop. Never dump passwords.
You can OPEN a live SSH console popup with open_console. os_exec, bmc_get, bmc_reset, get_health, and log reads also stream the command and response into that console so the operator can watch. Prefer open_console when troubleshooting a node. SSH is the right move for dmesg, nvidia-smi, and OS logs. Use Redfish for BMC sensors and power.
You can edit a test script for ONE run with save_test_script. That copy sticks to that instance. Do not rewrite the catalog.

When troubleshooting, check SOPs with list_sops / read_sop and give concrete suggestions from those docs plus live data.

When there is a clear next move, call offer_choices with 2-4 tap targets. Each choice is a real tool you would run if they tap it. Labels are short actions: "Restart BMC on GB300-Temp", "Run DCGM on the live node", "Pull dmesg from root OS". Skip offer_choices on a simple greeting or yes/no. Include hint when useful.
On a greeting, one short line. Do not call tools. Live cluster numbers are already in this prompt.

Talk like you speak. Short bubbles. One or two lines. Heh, hey, leave it to me. Concrete numbers from tools.
Never use em dashes or en dashes. Never write like a corporate assistant. No numbered essays.
When the user says things like "add 6 racks GB300 cluster name", set the cluster name to GB300 (or the name they gave) and create that many MaxLPS GB300 racks with labels GB300-01, GB300-02, …
When they paste a node, call add_node. Example: "add this node OS: 10.10.3.50 nvidia / nvidia BMC 10.10.17.177: root / 0penBmc" means os_ip 10.10.3.50, os_username nvidia, os_password nvidia, bmc_ip 10.10.17.177, bmc_user root, bmc_password 0penBmc. BMC MAC is optional when BMC IP is set. Pass hall_name when they name a rack; matching names share a rack, a new name is its own rack. You may pass the whole line as raw. Then they can KVM and SSH it from Provision. patch_node edits a node already in inventory.
Match test names from the catalog (mlperf-gpt3-175b, ib_write_bw, dcgm-diag, nccl-all_reduce, and so on).
Store BMC/OS logins they paste. Do not invent credentials. Do not dump passwords.
"""

TOOLS: list[dict[str, Any]] = [
    {
        "type": "function",
        "name": "get_cluster_status",
        "description": "Overview health tiles, cluster name, rack count. Not the Dynamic Power limiter.",
        "parameters": {"type": "object", "properties": {}, "additionalProperties": False},
    },
    {
        "type": "function",
        "name": "get_health",
        "description": "Live BMC health: leak, temperature (GPU), fans, PSU, per-node errors.",
        "parameters": {"type": "object", "properties": {}, "additionalProperties": False},
    },
    {
        "type": "function",
        "name": "get_configuration",
        "description": "Sanitized config and inventory: cluster name, racks, BMC IPs, probe, controller URL. No passwords.",
        "parameters": {"type": "object", "properties": {}, "additionalProperties": False},
    },
    {
        "type": "function",
        "name": "list_tests",
        "description": "Launchable tests: Argus catalog grouped by component (GPU, network, interconnect, system, …) plus local MLPerf.",
        "parameters": {
            "type": "object",
            "properties": {"component": {"type": "string", "description": "Optional filter: gpu, network, interconnect, pcie, memory, cpu, storage, thermal, system, mlperf"}},
        },
    },
    {
        "type": "function",
        "name": "list_test_reports",
        "description": "Argus rack test reports (pdf/xlsx/json) and recent results.",
        "parameters": {
            "type": "object",
            "properties": {"serial": {"type": "string"}, "name": {"type": "string"}},
        },
    },
    {
        "type": "function",
        "name": "read_test_logs",
        "description": "Argus node or rack test logs. Pass BMC MAC, or a report path from list_test_reports.",
        "parameters": {
            "type": "object",
            "properties": {"bmc_mac": {"type": "string"}, "path": {"type": "string"}},
        },
    },
    {
        "type": "function",
        "name": "run_argus_test",
        "description": "Launch an Argus catalog test (DCGM, GPU-BURN, L10 Diagnostic, NCCL, …) on a live rack.",
        "parameters": {
            "type": "object",
            "properties": {
                "test": {"type": "string"},
                "rack": {"type": "string", "description": "Rack serial or label"},
            },
            "required": ["test"],
        },
    },
    {
        "type": "function",
        "name": "list_sanity",
        "description": "Latest L11-style sanity matrix (FW/SW/BOM/PSU/voltage/SEL) and golden file status.",
        "parameters": {
            "type": "object",
            "properties": {"rack": {"type": "string", "description": "Rack label, serial, or a node hostname on that rack"}},
        },
    },
    {
        "type": "function",
        "name": "run_sanity_check",
        "description": "Run firmware, software, BOM, PSU, voltage, and SEL checks on a rack. Compares against golden BOM/FW/SW.",
        "parameters": {
            "type": "object",
            "properties": {
                "rack": {"type": "string", "description": "Rack label, serial, or a node hostname such as Vera Rubin"},
                "checks": {"type": "array", "items": {"type": "string"}, "description": "Subset: fw, sw, bom, psu, voltage, sel"},
            },
        },
    },
    {
        "type": "function",
        "name": "capture_golden",
        "description": "Capture golden BOM, firmware, and software from one node. Needed before BOM/FW/SW sanity can PASS.",
        "parameters": {
            "type": "object",
            "properties": {"rack": {"type": "string", "description": "Rack label, serial, or a node hostname"}, "node": {"type": "string", "description": "Node serial, hostname, or N1"}},
        },
    },
    {
        "type": "function",
        "name": "generate_l11_report",
        "description": "Build L11 rack or L10 node PDF/CSV/JSON from the latest sanity run.",
        "parameters": {
            "type": "object",
            "properties": {
                "level": {"type": "string", "enum": ["L11", "L10"]},
                "rack": {"type": "string"},
                "node": {"type": "string"},
            },
        },
    },
    {
        "type": "function",
        "name": "list_test_runs",
        "description": "Recent workload/test runs with status and detail.",
        "parameters": {
            "type": "object",
            "properties": {"limit": {"type": "integer", "minimum": 1, "maximum": 50}},
        },
    },
    {
        "type": "function",
        "name": "start_test_agent",
        "description": "Dispatch a minion that keeps working after this turn. kind: sanity, argus, or catalog. Returns minion name and briefing. Operator watches them in the AGENTS radio.",
        "parameters": {
            "type": "object",
            "properties": {
                "kind": {"type": "string", "enum": ["sanity", "argus", "catalog"], "description": "sanity = L11 FW/SW/BOM. argus = live Argus catalog test. catalog = local MLPerf/NCCL/IB instance."},
                "test": {"type": "string", "description": "Argus or catalog test id, e.g. dcgm-diag, gpu-burn"},
                "rack": {"type": "string", "description": "Rack label, serial, or a node hostname such as Vera Rubin. Switches the Testing rack menu."},
                "name": {"type": "string", "description": "Optional override for the minion's callsign"},
            },
        },
    },
    {
        "type": "function",
        "name": "list_agents",
        "description": "Mona's minion roster: callsign, mission, and what each one is doing.",
        "parameters": {"type": "object", "properties": {}, "additionalProperties": False},
    },
    {
        "type": "function",
        "name": "stop_agent",
        "description": "Recall a minion by id or callsign. They drop the job and you review what they had.",
        "parameters": {
            "type": "object",
            "properties": {"id": {"type": "string"}},
            "required": ["id"],
        },
    },
    {
        "type": "function",
        "name": "open_agent",
        "description": "Open the AGENTS radio on a minion by id so the operator can watch their chat.",
        "parameters": {
            "type": "object",
            "properties": {"id": {"type": "string"}},
            "required": ["id"],
        },
    },
    {
        "type": "function",
        "name": "read_logs",
        "description": "Tail logs. source: tests (workload details), provision (SOL), server (uvicorn/app logs).",
        "parameters": {
            "type": "object",
            "properties": {
                "source": {"type": "string", "enum": ["tests", "provision", "server"]},
                "query": {"type": "string", "description": "Optional name/id filter"},
                "limit": {"type": "integer", "minimum": 10, "maximum": 200},
            },
            "required": ["source"],
        },
    },
    {
        "type": "function",
        "name": "set_cluster_name",
        "description": "Rename the cluster (Overview kicker).",
        "parameters": {
            "type": "object",
            "properties": {"name": {"type": "string"}},
            "required": ["name"],
        },
    },
    {
        "type": "function",
        "name": "set_power_cap",
        "description": "Set the Dynamic Power limiter on /maxlps (PATCH /api/cluster/power). There is no /api/limiter route. Use this for limiter/power cap/budget/GPU percent. budget_kw is the cluster budget, max_rack_kw the per-rack cap (20 kW min), gpu_power_percent 10-100.",
        "parameters": {
            "type": "object",
            "properties": {
                "budget_kw": {"type": "number", "description": "Cluster limiter budget in kW"},
                "max_rack_kw": {"type": "number", "description": "Per-rack cap in kW, minimum 20"},
                "gpu_power_percent": {"type": "number", "description": "GPU power cap percent 10-100"},
            },
        },
    },
    {
        "type": "function",
        "name": "add_racks",
        "description": "Create MaxLPS GB300 inventory racks. Optionally set cluster name first.",
        "parameters": {
            "type": "object",
            "properties": {
                "count": {"type": "integer", "minimum": 1, "maximum": 18},
                "prefix": {"type": "string", "description": "Label prefix, default GB300"},
                "hall": {"type": "string"},
                "cluster_name": {"type": "string", "description": "If set, rename cluster first"},
            },
            "required": ["count"],
        },
    },
    {
        "type": "function",
        "name": "remove_rack",
        "description": "Remove a rack from Cluster, Dynamic Power, and Provision by label or id. The next Argus sync adds it back if the PXE still has it.",
        "parameters": {
            "type": "object",
            "properties": {"label": {"type": "string"}, "id": {"type": "string"}},
        },
    },
    {
        "type": "function",
        "name": "launch_test",
        "description": "Start a catalog test/workload. Optional script is saved on that run only.",
        "parameters": {
            "type": "object",
            "properties": {
                "name": {"type": "string", "description": "Catalog name or label"},
                "id": {"type": "string", "description": "Existing instance to start"},
                "script": {"type": "string", "description": "Instance-only script override"},
            },
            "required": ["name"],
        },
    },
    {
        "type": "function",
        "name": "stop_test",
        "description": "Stop a running workload by id or name.",
        "parameters": {
            "type": "object",
            "properties": {"id": {"type": "string"}, "name": {"type": "string"}},
        },
    },
    {
        "type": "function",
        "name": "sync_inventory",
        "description": "Sync MaxLPS inventory from the Argus cluster controller.",
        "parameters": {"type": "object", "properties": {}, "additionalProperties": False},
    },
    {
        "type": "function",
        "name": "list_sops",
        "description": "List uploaded SOP / reference documents for troubleshooting.",
        "parameters": {"type": "object", "properties": {}, "additionalProperties": False},
    },
    {
        "type": "function",
        "name": "read_sop",
        "description": "Read an SOP / reference document by id or title.",
        "parameters": {
            "type": "object",
            "properties": {"id": {"type": "string"}, "title": {"type": "string"}},
        },
    },
    {
        "type": "function",
        "name": "list_nodes",
        "description": "Connected compute nodes: BMC/OS IPs and whether logins are stored. No passwords.",
        "parameters": {"type": "object", "properties": {}, "additionalProperties": False},
    },
    {
        "type": "function",
        "name": "bmc_get",
        "description": "GET a Redfish path on an inventory BMC. Direct HTTPS only.",
        "parameters": {
            "type": "object",
            "properties": {
                "node": {"type": "string", "description": "BMC IP, hostname, or rack label"},
                "path": {"type": "string", "description": "Redfish path starting /redfish/v1"},
            },
            "required": ["node", "path"],
        },
    },
    {
        "type": "function",
        "name": "bmc_reset",
        "description": "BMC ComputerSystem.Reset on an inventory node. ResetType On, ForceOff, GracefulShutdown, ForceRestart, GracefulRestart.",
        "parameters": {
            "type": "object",
            "properties": {
                "node": {"type": "string"},
                "reset_type": {
                    "type": "string",
                    "enum": ["On", "ForceOff", "GracefulShutdown", "ForceRestart", "GracefulRestart"],
                },
            },
            "required": ["node", "reset_type"],
        },
    },
    {
        "type": "function",
        "name": "open_console",
        "description": "Open the SSH console popup on a connected node. Optional first command runs there. via=bmc SSHes the BMC IP with BMC login; via=os (default) SSHes the host OS.",
        "parameters": {
            "type": "object",
            "properties": {
                "node": {"type": "string"},
                "command": {"type": "string", "description": "Optional first command"},
                "via": {"type": "string", "enum": ["os", "bmc"]},
            },
            "required": ["node"],
        },
    },
    {
        "type": "function",
        "name": "os_exec",
        "description": "Run a shell command over SSH. Default is the host OS (os_ip + root). via=bmc uses the BMC IP and BMC login on port 22.",
        "parameters": {
            "type": "object",
            "properties": {
                "node": {"type": "string", "description": "OS IP, hostname, BMC IP, or rack label"},
                "command": {"type": "string"},
                "via": {"type": "string", "enum": ["os", "bmc"]},
            },
            "required": ["node", "command"],
        },
    },
    {
        "type": "function",
        "name": "get_test_script",
        "description": "Read the catalog default script, or the instance copy on a run.",
        "parameters": {
            "type": "object",
            "properties": {"name": {"type": "string"}, "id": {"type": "string", "description": "Workload/run id"}},
        },
    },
    {
        "type": "function",
        "name": "save_test_script",
        "description": "Save a modified test script onto ONE run instance. Does not change the catalog.",
        "parameters": {
            "type": "object",
            "properties": {
                "script": {"type": "string"},
                "name": {"type": "string", "description": "Catalog name if creating a new instance"},
                "id": {"type": "string", "description": "Existing run id"},
            },
            "required": ["script"],
        },
    },
    {
        "type": "function",
        "name": "open_page",
        "description": "Switch the operator UI to a page: cluster, overview, racks, floor, maxlps (Dynamic Power — limiter, power cap, watts), network, provision, testing, settings, users.",
        "parameters": {
            "type": "object",
            "properties": {"path": {"type": "string", "description": "Route or alias, e.g. /maxlps, limiter, dynamic power, /testing, racks"}},
            "required": ["path"],
        },
    },
    {
        "type": "function",
        "name": "ui_fill",
        "description": "Type into a labeled field on the current page. Set submit true to send the form.",
        "parameters": {
            "type": "object",
            "properties": {
                "target": {"type": "string", "description": "Label, placeholder, name, or aria-label"},
                "value": {"type": "string"},
                "submit": {"type": "boolean"},
            },
            "required": ["target", "value"],
        },
    },
    {
        "type": "function",
        "name": "ui_click",
        "description": "Click a labeled button, tab, or link on the current page.",
        "parameters": {
            "type": "object",
            "properties": {"target": {"type": "string"}},
            "required": ["target"],
        },
    },
    {
        "type": "function",
        "name": "highlight_ui",
        "description": "Point at something on the operator UI. Target a rack label (GB300-01), region (health, racks, floor, limiter/maxlps/dynamic power, testing, cluster-name), or a visible button/field. Limiter and power go to Dynamic Power, not cluster.",
        "parameters": {
            "type": "object",
            "properties": {
                "target": {"type": "string"},
                "label": {"type": "string", "description": "Short caption, e.g. this rack"},
                "action": {"type": "string", "enum": ["look", "add", "remove", "edit", "click"]},
            },
            "required": ["target"],
        },
    },
    {
        "type": "function",
        "name": "patch_rack",
        "description": "Update a MaxLPS rack row in the database (label, hall, serial, notes).",
        "parameters": {
            "type": "object",
            "properties": {
                "rack": {"type": "string", "description": "Rack label or id"},
                "label": {"type": "string"},
                "hall": {"type": "string"},
                "serial": {"type": "string"},
                "notes": {"type": "string"},
            },
            "required": ["rack"],
        },
    },
    {
        "type": "function",
        "name": "patch_node",
        "description": "Update a connected node's database row: BMC/OS IP and login, hostname, serial. Writes Provision and MaxLPS inventory. Does not dump passwords.",
        "parameters": {
            "type": "object",
            "properties": {
                "node": {"type": "string"},
                "bmc_ip": {"type": "string"},
                "bmc_mac": {"type": "string"},
                "bmc_user": {"type": "string"},
                "bmc_password": {"type": "string"},
                "os_ip": {"type": "string"},
                "os_username": {"type": "string"},
                "os_password": {"type": "string"},
                "hostname": {"type": "string"},
                "name": {"type": "string"},
                "serial": {"type": "string"},
                "hall_name": {"type": "string", "description": "Rack name. Matching names share a rack."},
                "rack": {"type": "string"},
            },
            "required": ["node"],
        },
    },
    {
        "type": "function",
        "name": "add_node",
        "description": "Add or update a Provision node from BMC/OS IP and login. Use when they say add this node and paste OS/BMC addresses. MAC optional if BMC IP is set. hall_name groups nodes onto the same rack; omit it and the node is its own rack. Example raw: OS: 10.10.3.50 nvidia / nvidia BMC 10.10.17.177: root / 0penBmc",
        "parameters": {
            "type": "object",
            "properties": {
                "bmc_ip": {"type": "string"},
                "bmc_mac": {"type": "string"},
                "bmc_user": {"type": "string"},
                "bmc_password": {"type": "string"},
                "os_ip": {"type": "string"},
                "os_username": {"type": "string"},
                "os_password": {"type": "string"},
                "serial": {"type": "string"},
                "name": {"type": "string"},
                "hall_name": {"type": "string", "description": "Rack name. Same name shares a rack."},
                "rack": {"type": "string"},
                "raw": {"type": "string", "description": "Full pasted line to parse if fields are not split out"},
            },
        },
    },
    {
        "type": "function",
        "name": "offer_choices",
        "description": "Show tap-to-run next steps. Call this at the end of every reply.",
        "parameters": {
            "type": "object",
            "properties": {
                "choices": {
                    "type": "array",
                    "minItems": 1,
                    "maxItems": 4,
                    "items": {
                        "type": "object",
                        "properties": {
                            "label": {"type": "string"},
                            "tool": {"type": "string"},
                            "arguments": {"type": "object"},
                            "hint": {"type": "string"},
                        },
                        "required": ["label", "tool"],
                    },
                }
            },
            "required": ["choices"],
        },
    },
    {
        "type": "function",
        "name": "list_attachments",
        "description": "Files the operator attached in chat (this session and recent).",
        "parameters": {"type": "object", "properties": {}, "additionalProperties": False},
    },
    {
        "type": "function",
        "name": "read_attachment",
        "description": "Read an attached file by id or filename. Images return a short caption of stored metadata; text/pdf/xlsx return extracted text.",
        "parameters": {
            "type": "object",
            "properties": {"id": {"type": "string"}, "filename": {"type": "string"}},
        },
    },
    {
        "type": "function",
        "name": "list_app_routes",
        "description": "List L12 HTTP API routes you can call with call_app (method + path).",
        "parameters": {
            "type": "object",
            "properties": {
                "query": {"type": "string", "description": "Optional filter, e.g. provision or maxlps"},
            },
        },
    },
    {
        "type": "function",
        "name": "call_app",
        "description": "Call any L12 /api route as the signed-in operator. Same auth as the UI. Do not call /api/agent/chat.",
        "parameters": {
            "type": "object",
            "properties": {
                "method": {"type": "string", "enum": ["GET", "POST", "PUT", "PATCH", "DELETE"]},
                "path": {"type": "string", "description": "Path starting /api/, e.g. /api/cluster/overview"},
                "query": {"type": "object", "description": "Query string fields"},
                "body": {"type": "object", "description": "JSON body for POST, PUT, PATCH"},
            },
            "required": ["method", "path"],
        },
    },
]

_KEY_OVERRIDE: str | None = None
_GROK_AUTH = Path.home() / ".grok" / "auth.json"


def _grok_cli_token() -> str:
    """Session token from `grok login` on this machine. Not a console xai- key."""
    try:
        data = json.loads(_GROK_AUTH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError, UnicodeError):
        return ""
    if not isinstance(data, dict):
        return ""
    from datetime import datetime, timezone

    now = datetime.now(timezone.utc)
    best = ""
    best_exp = ""
    for entry in data.values():
        if not isinstance(entry, dict):
            continue
        tok = (entry.get("key") or "").strip()
        if not tok:
            continue
        exp_raw = str(entry.get("expires_at") or "")
        if exp_raw:
            try:
                exp = datetime.fromisoformat(exp_raw.replace("Z", "+00:00"))
                if exp.tzinfo is None:
                    exp = exp.replace(tzinfo=timezone.utc)
                if exp <= now:
                    continue
            except ValueError:
                pass
        if not best or exp_raw > best_exp:
            best = tok
            best_exp = exp_raw
    return best


def current_key(db: Session | None = None) -> str:
    if _KEY_OVERRIDE:
        return _KEY_OVERRIDE
    grok = _grok_cli_token()
    if grok:
        return grok
    if db is not None:
        from .models import ClusterSettings

        row = db.get(ClusterSettings, 1)
        stored = (getattr(row, "xai_api_key", None) or "").strip() if row else ""
        if stored:
            return stored
    return (settings.xai_api_key or "").strip()


def enabled(db: Session | None = None) -> bool:
    return bool(current_key(db))


def key_hint(db: Session | None = None) -> str | None:
    key = current_key(db)
    if not key:
        return None
    if _grok_cli_token() and key == _grok_cli_token():
        return "login"
    return key[-4:] if len(key) >= 4 else "set"


def set_api_key(db: Session, key: str | None) -> dict[str, Any]:
    global _KEY_OVERRIDE
    from .models import ClusterSettings

    cleaned = (key or "").strip() or None
    _KEY_OVERRIDE = cleaned
    row = db.get(ClusterSettings, 1)
    if row is None:
        row = ClusterSettings(id=1, name="Firmus", xai_api_key=cleaned)
        db.add(row)
    else:
        row.xai_api_key = cleaned
    db.commit()
    return {"configured": bool(current_key(db)), "hint": key_hint(db)}


def _words(text: str) -> set[str]:
    out: set[str] = set()
    buf: list[str] = []
    for ch in (text or "").lower():
        if ch.isalnum():
            buf.append(ch)
        elif buf:
            out.add("".join(buf))
            buf = []
    if buf:
        out.add("".join(buf))
    return out


def _needs_reasoning(
    text: str,
    *,
    attachments: bool = False,
    execute: dict[str, Any] | None = None,
) -> bool:
    if attachments:
        return True
    tool = str((execute or {}).get("tool") or "")
    if tool in _DEEP_TOOLS:
        return True
    raw = (text or "").strip()
    if len(raw) > 500:
        return True
    words = _words(raw)
    if words & _ANALYZE_WORDS:
        return True
    if words & {"look", "read", "open", "opened", "check"} and words & _EVIDENCE_WORDS:
        return True
    return False


def operator_model(
    task: str = "",
    *,
    attachments: bool = False,
    execute: dict[str, Any] | None = None,
) -> str:
    if _needs_reasoning(task, attachments=attachments, execute=execute):
        return _SMART_MODEL
    return _FAST_MODEL


def _is_smalltalk(text: str) -> bool:
    raw = (text or "").strip().lower()
    if not raw or len(raw) > 48:
        return False
    cleaned = "".join(ch if ch.isalnum() or ch.isspace() else " " for ch in raw)
    cleaned = " ".join(cleaned.split())
    return cleaned in _SMALLTALK


def _brief_cluster(db: Session) -> str:
    try:
        st = tool_get_cluster_status(db)
    except Exception:
        return ""
    racks = [r for r in (st.get("racks") or []) if isinstance(r, dict)]
    labels = ", ".join(str(r.get("label") or r.get("id") or "") for r in racks[:8] if r.get("label") or r.get("id"))
    health = st.get("health") or {}
    live = ""
    try:
        from . import agent_jobs

        live = agent_jobs.running_summary()
    except Exception:
        live = ""
    agents = f" Minions out: {live}." if live else " No minions out."
    return (
        f"\nLive snapshot: cluster {st.get('name')}; {len(racks)} racks"
        f"{f' ({labels})' if labels else ''}; leak={health.get('leak')}; "
        f"temp={health.get('temperature')}; fans={health.get('fans')}; psu={health.get('psu')}."
        f"{agents} "
        "Use tools only if you need a deeper read."
    )


def _dump(obj: Any, limit: int = 8000) -> str:
    text = json.dumps(obj, default=str, ensure_ascii=False)
    return text if len(text) <= limit else text[: limit - 20] + "…[truncated]"


def _pretty(obj: Any, limit: int = 6000) -> str:
    if obj is None:
        return ""
    if isinstance(obj, str):
        return obj if len(obj) <= limit else obj[: limit - 20] + "…[truncated]"
    text = json.dumps(obj, indent=2, default=str, ensure_ascii=False)
    return text if len(text) <= limit else text[: limit - 20] + "…[truncated]"


_CONSOLE_TOOLS = {
    "bmc_get",
    "bmc_reset",
    "os_exec",
    "open_console",
    "get_health",
    "read_logs",
    "read_test_logs",
}


def _console_command(name: str, args: dict[str, Any], data: dict[str, Any]) -> str:
    if name == "bmc_get":
        path = str(args.get("path") or data.get("path") or "/redfish/v1")
        ip = str(data.get("bmc_ip") or args.get("node") or "")
        return f"GET {ip}{path}" if ip else f"GET {path}"
    if name == "bmc_reset":
        kind = str(args.get("reset_type") or data.get("reset_type") or "")
        ip = str(data.get("bmc_ip") or args.get("node") or "")
        return f"POST {ip} ComputerSystem.Reset ResetType={kind}".strip()
    if name in ("os_exec", "open_console"):
        return str(args.get("command") or data.get("command") or "ssh")
    if name == "read_test_logs":
        return f"read_logs {args.get('bmc_mac') or args.get('path') or ''}".strip()
    if name == "read_logs":
        return f"read_logs {args.get('source') or 'server'}"
    if name == "get_health":
        return "get_health"
    return name


def _make_console_ui(name: str, args: dict[str, Any], result: Any) -> dict[str, Any] | None:
    if name not in _CONSOLE_TOOLS:
        return None
    data = result if isinstance(result, dict) else {"result": result}
    node = str(data.get("bmc_ip") or data.get("host") or args.get("node") or ("cluster" if name == "get_health" else name))
    host = str(data.get("bmc_ip") or data.get("host") or args.get("node") or "")
    user = str(data.get("user") or "root")
    err = data.get("error")
    ok = not (data.get("ok") is False or err)
    stdout = data.get("stdout") if name in ("os_exec", "open_console") else None
    if not stdout:
        body = {k: v for k, v in data.items() if k not in ("ui", "stdout", "stderr")}
        stdout = _pretty(body)
    return {
        "kind": "console",
        "node": node,
        "host": host,
        "user": user,
        "command": _console_command(name, args, data),
        "stdout": stdout or "",
        "stderr": str(data.get("stderr") or ""),
        "exit": data.get("exit"),
        "error": err,
        "ok": ok,
    }


def _attach_console_ui(name: str, args: dict[str, Any], result: Any) -> Any:
    if not isinstance(result, dict):
        ui = _make_console_ui(name, args, result)
        return {"result": result, "ui": ui} if ui else result
    existing = result.get("ui")
    if isinstance(existing, dict) and existing.get("kind") == "console":
        return result
    if isinstance(existing, list) and any(isinstance(x, dict) and x.get("kind") == "console" for x in existing):
        return result
    ui = _make_console_ui(name, args, result)
    if not ui:
        return result
    out = dict(result)
    out["ui"] = ui
    return out


def _need_edit(user: User) -> None:
    if not has_rank(user.role, EDITOR_FLOOR):
        raise PermissionError("Requires Manager or above to change the cluster")


def _tail_file(path: Path, n: int) -> str:
    if not path.is_file():
        return ""
    try:
        data = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        return f"(unreadable: {exc})"
    lines = data.splitlines()
    return "\n".join(lines[-n:])


_TEXT_EXT = {
    ".txt",
    ".md",
    ".csv",
    ".log",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
    ".py",
    ".sh",
    ".conf",
    ".ini",
    ".toml",
    ".cfg",
}
_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
_DOC_EXT = {".pdf", ".docx", ".xlsx"}
ATTACH_EXT = _TEXT_EXT | _IMAGE_EXT | _DOC_EXT
_TEXT_LIMIT = 80_000


def extract_upload(filename: str, data: bytes) -> dict[str, str]:
    """Turn uploaded bytes into text (and a kind) Mona can read."""
    ext = Path(filename or "file").suffix.lower()
    mime = "application/octet-stream"
    if ext in _IMAGE_EXT:
        mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp", ".gif": "image/gif"}.get(
            ext, "image/jpeg"
        )
        return {"kind": "image", "text": "", "mime": mime}
    if ext in _TEXT_EXT:
        text = data.decode("utf-8", errors="replace")
        if ext == ".json":
            try:
                text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
            except json.JSONDecodeError:
                pass
        mime = "text/plain"
        if len(text) > _TEXT_LIMIT:
            text = text[:_TEXT_LIMIT] + "\n[truncated]"
        return {"kind": "text", "text": text, "mime": mime}
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:40])
        if len(text) > _TEXT_LIMIT:
            text = text[:_TEXT_LIMIT] + "\n[truncated]"
        return {"kind": "text", "text": text, "mime": "application/pdf"}
    if ext == ".docx":
        from docx import Document

        doc = Document(BytesIO(data))
        text = "\n".join(p.text for p in doc.paragraphs)
        if len(text) > _TEXT_LIMIT:
            text = text[:_TEXT_LIMIT] + "\n[truncated]"
        return {"kind": "text", "text": text, "mime": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
    if ext == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        chunks: list[str] = []
        for sheet in wb.worksheets[:8]:
            chunks.append(f"# {sheet.title}")
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= 80:
                    chunks.append("…")
                    break
                cells = ["" if c is None else str(c) for c in row]
                if any(cells):
                    chunks.append("\t".join(cells))
        text = "\n".join(chunks)
        if len(text) > _TEXT_LIMIT:
            text = text[:_TEXT_LIMIT] + "\n[truncated]"
        return {"kind": "text", "text": text, "mime": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
    return {"kind": "binary", "text": "", "mime": mime}


def save_attachment(db: Session, user: User, filename: str, data: bytes, mime: str | None = None) -> AgentAttachment:
    from .storage import store_bytes

    name = Path(filename or "file").name
    ext = Path(name).suffix.lower() or ".bin"
    parsed = extract_upload(name, data)
    url = store_bytes(data, ext if ext.startswith(".") else f".{ext}", parsed.get("mime") or mime)
    row = AgentAttachment(
        user_id=user.id,
        filename=name[:255],
        mime=(mime or parsed.get("mime") or "application/octet-stream")[:120],
        kind=parsed["kind"],
        text=parsed.get("text") or "",
        url=url,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _attachment_public(row: AgentAttachment) -> dict[str, Any]:
    return {
        "id": row.id,
        "filename": row.filename,
        "kind": row.kind,
        "mime": row.mime,
        "chars": len(row.text or ""),
        "url": row.url,
        "created_at": row.created_at,
    }


def _user_attachments(db: Session, user: User, ids: list[str] | None = None) -> list[AgentAttachment]:
    from sqlalchemy import select

    if ids:
        want = [i for i in ids if i]
        rows = list(db.scalars(select(AgentAttachment).where(AgentAttachment.id.in_(want))).all())
        by_id = {r.id: r for r in rows if r.user_id == user.id}
        return [by_id[i] for i in want if i in by_id]
    return list(
        db.scalars(
            select(AgentAttachment).where(AgentAttachment.user_id == user.id).order_by(AgentAttachment.created_at.desc()).limit(30)
        ).all()
    )


def tool_list_attachments(db: Session, user: User, **_kw: Any) -> dict[str, Any]:
    rows = _user_attachments(db, user)
    return {"files": [_attachment_public(r) for r in rows]}


def tool_read_attachment(db: Session, user: User, id: str | None = None, filename: str | None = None, **_kw: Any) -> dict[str, Any]:
    from sqlalchemy import select

    row = db.get(AgentAttachment, id) if id else None
    if row is not None and row.user_id != user.id:
        row = None
    if row is None and filename:
        needle = filename.strip().lower()
        rows = list(db.scalars(select(AgentAttachment).where(AgentAttachment.user_id == user.id)).all())
        row = next((r for r in rows if r.filename.lower() == needle or r.id.lower() == needle), None)
        if row is None:
            row = next((r for r in rows if needle in r.filename.lower()), None)
    if row is None:
        return {"ok": False, "error": "attachment not found", **tool_list_attachments(db, user)}
    text = row.text or ""
    if row.kind == "image" and not text:
        text = f"(image {row.filename}. Vision already saw it if it was attached to this turn.)"
    if len(text) > 12000:
        text = text[:12000] + "\n[truncated]"
    return {"ok": True, **_attachment_public(row), "text": text}


def _jpeg_for_model(data: bytes) -> bytes:
    from PIL import Image

    im = Image.open(BytesIO(data))
    if im.mode != "RGB":
        im = im.convert("RGB")
    im.thumbnail((1600, 1600))
    buf = BytesIO()
    im.save(buf, format="JPEG", quality=72, optimize=True)
    return buf.getvalue()


def _image_part(row: AgentAttachment) -> dict[str, Any] | None:
    from .storage import load_bytes

    raw = load_bytes(row.url)
    if not raw:
        return None
    try:
        jpeg = _jpeg_for_model(raw)
    except Exception:
        return None
    b64 = base64.b64encode(jpeg).decode("ascii")
    return {"type": "input_image", "image_url": f"data:image/jpeg;base64,{b64}", "detail": "high"}


_SECRET_KEY_BITS = ("password", "secret", "token", "authorization", "xai_api_key", "api_key", "jwt_secret")


def _redact(obj: Any, depth: int = 0) -> Any:
    if depth > 8:
        return obj
    if isinstance(obj, dict):
        out = {}
        for key, val in obj.items():
            name = str(key).lower()
            if name in _SECRET_KEY_BITS or "password" in name:
                out[key] = "***" if val else val
            else:
                out[key] = _redact(val, depth + 1)
        return out
    if isinstance(obj, list):
        return [_redact(item, depth + 1) for item in obj[:200]]
    return obj


_BLOCKED_API = {"/api/agent/chat"}


def _norm_api_path(raw: str) -> str:
    s = (raw or "").strip().split("?")[0].split("#")[0]
    if not s or "://" in s or ".." in s or "\\" in s:
        return ""
    if not s.startswith("/"):
        s = "/" + s
    if not s.startswith("/api/"):
        s = "/api" + s
    s = "/" + "/".join(p for p in s.split("/") if p)
    if not s.startswith("/api/") or s in _BLOCKED_API or s.startswith("/api/agent/chat"):
        return ""
    return s


def tool_list_app_routes(query: str | None = None, **_kw: Any) -> dict[str, Any]:
    from .main import app

    needle = (query or "").strip().lower()
    routes: list[dict[str, Any]] = []
    for route in app.routes:
        path = getattr(route, "path", None)
        methods = getattr(route, "methods", None)
        if not path or not methods:
            continue
        if not str(path).startswith("/api/") or str(path).startswith("/api/agent/chat"):
            continue
        verbs = sorted(m for m in methods if m in {"GET", "POST", "PUT", "PATCH", "DELETE"})
        if not verbs:
            continue
        blob = f"{' '.join(verbs)} {path} {getattr(route, 'name', '') or ''}".lower()
        if needle and needle not in blob:
            continue
        item: dict[str, Any] = {"methods": verbs, "path": path}
        summary = getattr(route, "summary", None) or getattr(route, "name", None)
        if summary:
            item["name"] = str(summary)
        routes.append(item)
    routes.sort(key=lambda r: r["path"])
    return {"count": len(routes), "routes": routes[:180]}


def tool_call_app(
    method: str,
    path: str,
    query: dict[str, Any] | None = None,
    body: dict[str, Any] | None = None,
    **_kw: Any,
) -> dict[str, Any]:
    from starlette.testclient import TestClient

    from .main import app

    verb = (method or "GET").strip().upper()
    if verb not in {"GET", "POST", "PUT", "PATCH", "DELETE"}:
        return {"ok": False, "error": "method must be GET, POST, PUT, PATCH, or DELETE"}
    dest = _norm_api_path(path)
    if not dest:
        return {"ok": False, "error": "path must be an /api/ route (not /api/agent/chat)"}
    auth = _AUTH_HEADER.get() or ""
    headers = {"Authorization": auth} if auth else {}
    params = {str(k): v for k, v in (query or {}).items() if v is not None} if isinstance(query, dict) else None
    payload = body if verb in {"POST", "PUT", "PATCH"} and isinstance(body, dict) else None
    try:
        client = TestClient(app, raise_server_exceptions=False, lifespan="off")
        resp = client.request(verb, dest, headers=headers, params=params, json=payload)
    except TypeError:
        client = TestClient(app, raise_server_exceptions=False)
        resp = client.request(verb, dest, headers=headers, params=params, json=payload)
    except Exception as exc:
        return {"ok": False, "method": verb, "path": dest, "error": str(exc)[:240]}
    parsed: Any
    try:
        parsed = resp.json()
    except Exception:
        parsed = (resp.text or "")[:4000]
    parsed = _redact(parsed)
    dumped = _dump(parsed, 6000)
    try:
        parsed = json.loads(dumped) if dumped.startswith("{") or dumped.startswith("[") else dumped
    except json.JSONDecodeError:
        parsed = dumped
    return {
        "ok": 200 <= resp.status_code < 400,
        "status": resp.status_code,
        "method": verb,
        "path": dest,
        "body": parsed,
    }


def tool_get_cluster_status(db: Session, **_kw: Any) -> dict[str, Any]:
    from . import bmc_health
    from .routers.cluster import cluster_name

    health = bmc_health.snapshot()
    racks = []
    try:
        from . import maxlps_inventory as inv

        racks = [{"id": r.id, "label": r.label, "hall": r.hall, "nodes": sum(1 for n in r.nodes if (n.bmc_ip or "").strip())} for r in inv.load_racks(db)]
    except Exception as exc:
        racks = [{"error": str(exc)}]
    return {
        "name": cluster_name(db),
        "racks": racks,
        "ui": _point_ui("look", "overview", "cluster status", path="/cluster?tab=overview"),
        "health": {
            "leak": health.get("leak"),
            "temperature": health.get("temperature"),
            "system": health.get("system"),
            "fans": health.get("fans"),
            "psu": health.get("psu"),
            "power_shelf": health.get("power_shelf"),
            "node_count": health.get("node_count"),
            "error": health.get("error"),
            "poll_age_s": health.get("poll_age_s"),
        },
    }


def tool_get_health(**_kw: Any) -> dict[str, Any]:
    from . import bmc_health

    snap = bmc_health.snapshot()
    nodes = []
    for n in snap.get("nodes") or []:
        nodes.append(
            {
                "bmc_ip": n.get("bmc_ip"),
                "status": n.get("status"),
                "error": n.get("error"),
                "product": n.get("product"),
                "thermal": {
                    "hottest": (n.get("thermal") or {}).get("hottest"),
                    "hottest_c": (n.get("thermal") or {}).get("hottest_c"),
                    "gpu_temps": (n.get("thermal") or {}).get("gpu_temps"),
                },
                "leak": (n.get("leak") or {}).get("status"),
                "fans": {"running": (n.get("fans") or {}).get("running"), "count": (n.get("fans") or {}).get("count")},
                "power_w": (n.get("power") or {}).get("watts"),
            }
        )
    return {
        "leak": snap.get("leak"),
        "temperature": snap.get("temperature"),
        "system": snap.get("system"),
        "fans": snap.get("fans"),
        "psu": snap.get("psu"),
        "power_shelf": snap.get("power_shelf"),
        "nodes": nodes,
        "error": snap.get("error"),
        "via": snap.get("via"),
        "ui": _point_ui("look", "health", "sniffing health", path="/cluster?tab=overview"),
    }


def tool_get_configuration(db: Session, **_kw: Any) -> dict[str, Any]:
    from . import maxlps_inventory as inv
    from .routers.cluster import cluster_name

    racks = []
    for r in inv.load_racks(db):
        d = inv.dump_rack(r, secrets=False)
        racks.append(
            {
                "id": d["id"],
                "label": d["label"],
                "serial": d.get("serial"),
                "hall": d.get("hall"),
                "model": d.get("model"),
                "origin": d.get("origin"),
                "nodes": [
                    {"index": n["index"], "hostname": n.get("hostname"), "bmc_ip": n.get("bmc_ip"), "bmc_user": n.get("bmc_user")}
                    for n in d.get("nodes") or []
                    if (n.get("bmc_ip") or n.get("hostname"))
                ],
            }
        )
    return {
        "cluster_name": cluster_name(db),
        "bmc_probe_ip": settings.bmc_probe_ip,
        "cluster_controller_url": settings.cluster_controller_url,
        "cluster_pxe_host": settings.cluster_pxe_host,
        "racks": racks,
    }


def tool_list_tests(component: str | None = None, **_kw: Any) -> dict[str, Any]:
    from . import argus_tests, sanity

    cat = argus_tests.catalog()
    tests = list(cat.get("tests") or [])
    if not tests:
        tests = [{**t, "component": t["kind"], "source": "local", "id": t["name"]} for t in TEST_CATALOG]
    for spec in sanity.script_catalog().get("scripts") or []:
        tests.append(
            {
                "id": spec["id"],
                "name": spec["id"],
                "label": spec["label"],
                "component": spec.get("component"),
                "kind": spec.get("component"),
                "detail": spec.get("detail"),
                "source": "l11",
            }
        )
    needle = (component or "").strip().lower()
    if needle:
        tests = [t for t in tests if str(t.get("component") or t.get("kind") or "").lower() == needle]
    return {
        "source": cat.get("source") or "l11",
        "components": cat.get("components") or [],
        "tests": tests,
        "error": cat.get("error"),
        "ui": _point_ui("look", "catalog", "browsing tests", path="/testing"),
    }


def tool_list_test_reports(serial: str | None = None, name: str | None = None, **_kw: Any) -> dict[str, Any]:
    from . import argus_tests, l11_report

    bundle = argus_tests.list_reports(serial, name)
    local = l11_report.list_reports()
    runs = argus_tests.active_runs()
    rows = list(local.get("reports") or []) + list(bundle.get("reports") or [])
    return {
        "reports": rows[:40],
        "active": runs,
        "count": len(rows),
        "error": bundle.get("error"),
        "ui": _point_ui("look", "reports", "digging reports", path="/testing"),
    }


def tool_read_test_logs(bmc_mac: str | None = None, path: str | None = None, **_kw: Any) -> dict[str, Any]:
    from . import argus_tests

    try:
        if path:
            return {"ok": True, "path": path, "body": argus_tests.read_log(path)}
        if bmc_mac:
            return {"ok": True, "bmc_mac": bmc_mac, "body": argus_tests.node_logs(bmc_mac)}
        return {"ok": True, "body": argus_tests.logs_index()}
    except Exception as exc:
        return {"ok": False, "error": str(exc)[:240]}


def _switch_testing_rack(db: Session, rack: str | None, *, pane: str = "Sanity") -> list[dict[str, Any]]:
    """Point the Testing page at this rack. A node hostname selects that node's rack."""
    from . import sanity

    ui: list[dict[str, Any]] = [{"kind": "navigate", "path": "/testing"}]
    named = (rack or "").strip()
    chosen = sanity.pick_rack(db, named, named) if named else None
    if chosen is not None:
        ui.append(
            {
                "kind": "testing-rack",
                "rack": chosen.label,
                "serial": chosen.serial or chosen.id,
                "id": chosen.id,
                "label": chosen.label,
            }
        )
    if pane:
        ui.append({"kind": "click", "target": pane})
    return ui


def tool_run_argus_test(db: Session, user: User, test: str, rack: str | None = None, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from . import argus_tests, sanity

    name = (test or "").strip()
    if not name:
        return {"ok": False, "error": "give a test id"}
    named = (rack or "").strip()
    chosen = sanity.pick_rack(db, named or None, named or None)
    if named and chosen is None:
        return {"ok": False, "error": f"no rack matches {named}"}
    if chosen is None or not (chosen.serial or "").strip():
        return {"ok": False, "error": "no Argus rack to run on"}
    nodes = []
    for n in chosen.nodes:
        mac = (n.bmc_mac or "").strip()
        if mac:
            nodes.append({"bmc_mac": mac, "bmc_ip": n.bmc_ip, "hostname": n.hostname, "location_index": str(n.index)})
    try:
        result = argus_tests.launch(
            [name],
            [{"rack_id": chosen.serial, "rack_name": chosen.label, "nodes": nodes}],
        )
    except Exception as exc:
        return {"ok": False, "error": str(exc)[:240]}
    return {
        "ok": True,
        "test": name,
        "rack": chosen.label,
        "serial": chosen.serial,
        "result": result,
        "ui": _switch_testing_rack(db, chosen.label, pane="Tests")
        + _point_ui("look", "catalog", f"kicking {name}", path="/testing"),
    }


def tool_list_sanity(db: Session, rack: str | None = None, **_kw: Any) -> dict[str, Any]:
    from . import sanity

    data = sanity.load_status(rack)
    nodes = data.get("nodes") or {}
    slim = {
        k: {
            "serial": v.get("serial"),
            "overall": v.get("overall"),
            "checks": {cid: (chk or {}).get("status") for cid, chk in (v.get("checks") or {}).items()},
        }
        for k, v in nodes.items()
    }
    return {
        "overall": data.get("overall"),
        "status": data.get("status"),
        "rack": data.get("rack_label") or data.get("rack_serial"),
        "golden": data.get("golden"),
        "nodes": slim,
        "job": data.get("job"),
        "ui": _switch_testing_rack(db, rack)
        + _point_ui("look", "sanity", "sanity matrix", path="/testing"),
    }


def tool_run_sanity_check(db: Session, user: User, rack: str | None = None, checks: list[str] | None = None, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from . import sanity

    out = sanity.start_sanity(db, rack_serial=rack, rack_name=rack, checks=checks)
    shown = str(out.get("rack_label") or rack or "")
    out["ui"] = _switch_testing_rack(db, shown) + _point_ui("look", "sanity", "running sanity", path="/testing")
    return out


def tool_capture_golden(db: Session, user: User, rack: str | None = None, node: str | None = None, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from . import sanity

    out = sanity.capture_golden(db, rack_serial=rack, rack_name=rack, node=node)
    slim = {k: v for k, v in out.items() if k != "bom"}
    shown = str(out.get("rack_label") or rack or "")
    slim["ui"] = _switch_testing_rack(db, shown) + _point_ui("look", "golden", "capturing golden", path="/testing")
    return slim


def tool_generate_l11_report(user: User, level: str | None = None, rack: str | None = None, node: str | None = None, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from . import l11_report

    try:
        out = l11_report.generate(level=level or "L11", rack_serial=rack, node_serial=node)
    except (FileNotFoundError, ValueError) as exc:
        return {"ok": False, "error": str(exc), "ui": _point_ui("look", "reports", "need a sanity run first", path="/testing")}
    out["ui"] = _point_ui("look", "reports", f"built {out.get('level')} report", path="/testing")
    return out


def tool_start_test_agent(
    db: Session,
    user: User,
    kind: str | None = None,
    test: str | None = None,
    rack: str | None = None,
    name: str | None = None,
    **_kw: Any,
) -> dict[str, Any]:
    _need_edit(user)
    from . import agent_jobs, sanity

    named = (rack or "").strip()
    chosen = sanity.pick_rack(db, named or None, named or None) if named else None
    if named and chosen is None:
        return {"ok": False, "error": f"no rack matches {named}"}
    rack_tag = chosen.label if chosen is not None else rack
    out = agent_jobs.start(db, user, kind=kind, test=test, rack=rack_tag, name=name)
    if out.get("ok"):
        out["rack"] = chosen.label if chosen is not None else rack
        out["serial"] = (chosen.serial or chosen.id) if chosen is not None else None
        out["ui"] = _switch_testing_rack(db, rack_tag)
    return out


def tool_list_agents(**_kw: Any) -> dict[str, Any]:
    from . import agent_jobs

    rows = agent_jobs.list_jobs()
    live = [r for r in rows if r.get("status") == "running"]
    return {
        "ok": True,
        "running": len(live),
        "agents": rows,
        "ui": [{"kind": "agents"}],
    }


def tool_stop_agent(db: Session, user: User, id: str, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from . import agent_jobs

    out = agent_jobs.stop(id, db, user)
    out["ui"] = [{"kind": "agents"}]
    return out


def tool_open_agent(id: str | None = None, **_kw: Any) -> dict[str, Any]:
    from . import agent_jobs

    jid = (id or "").strip()
    row = agent_jobs.public(jid) if jid else None
    return {
        "ok": True,
        "id": jid or None,
        "minion": (row or {}).get("minion"),
        "mission": (row or {}).get("mission"),
        "ui": [{"kind": "agents", "id": jid}] if jid else [{"kind": "agents"}],
    }


def tool_list_test_runs(db: Session, limit: int = 20, **_kw: Any) -> dict[str, Any]:
    from sqlalchemy import select

    rows = list(db.scalars(select(Workload).order_by(Workload.created_at.desc()).limit(int(limit or 20))).all())
    return {
        "runs": [
            {
                "id": w.id,
                "name": w.name,
                "kind": w.kind,
                "status": w.status,
                "detail": w.detail,
                "gpu_request": w.gpu_request,
                "pods_requested": w.pods_requested,
                "started_at": w.started_at,
                "finished_at": w.finished_at,
            }
            for w in rows
        ]
    }


def tool_read_logs(db: Session, source: str = "server", query: str | None = None, limit: int = 80, **_kw: Any) -> dict[str, Any]:
    n = max(10, min(int(limit or _LOG_TAIL), 200))
    q = (query or "").strip().lower()
    if source == "tests":
        runs = tool_list_test_runs(db, limit=n)["runs"]
        if q:
            runs = [r for r in runs if q in f"{r['name']} {r['id']} {r.get('detail') or ''}".lower()]
        return {"source": "tests", "entries": runs}
    if source == "provision":
        from sqlalchemy import select

        from .models import InventoryNode

        nodes = list(db.scalars(select(InventoryNode)).all())
        out = []
        for node in nodes:
            blob = node.sol_log or ""
            name = node.name or node.serial or node.id
            if q and q not in f"{name} {blob}".lower():
                continue
            lines = blob.splitlines()[-n:]
            out.append({"id": node.id, "name": name, "status": node.provision_status, "log": "\n".join(lines)})
        return {"source": "provision", "entries": out[:20]}
    root = Path(__file__).resolve().parents[2]
    files = [
        root / "backend" / "uvicorn-8001.log",
        root / "backend" / "uvicorn-8001.out",
        root / "backend" / "uvicorn-8001.err",
        root / "frontend" / "next-3001.out",
        root / "frontend" / "next-3001.err",
        root / ".local" / "test-logs",
    ]
    chunks = []
    for p in files:
        if p.is_dir():
            kids = sorted(p.glob("**/*"), key=lambda x: x.stat().st_mtime if x.is_file() else 0, reverse=True)[:8]
            for k in kids:
                if k.is_file() and k.stat().st_size < 2_000_000:
                    text = _tail_file(k, n)
                    if q and q not in text.lower() and q not in k.name.lower():
                        continue
                    if text:
                        chunks.append({"file": str(k.relative_to(root)), "text": text})
            continue
        text = _tail_file(p, n)
        if not text:
            continue
        if q and q not in text.lower():
            continue
        chunks.append({"file": p.name, "text": text})
    return {"source": "server", "entries": chunks}


def tool_set_cluster_name(db: Session, user: User, name: str, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from fastapi import HTTPException

    from .routers.cluster import set_cluster_name

    try:
        saved = set_cluster_name(db, name)
        return {
            "ok": True,
            "name": saved,
            "ui": _point_ui("edit", "cluster-name", f"cluster is {saved}", path="/cluster"),
        }
    except HTTPException as exc:
        return {"ok": False, "error": str(exc.detail)}


def tool_set_power_cap(
    db: Session,
    user: User,
    budget_kw: float | None = None,
    max_rack_kw: float | None = None,
    gpu_power_percent: float | None = None,
    **_kw: Any,
) -> dict[str, Any]:
    _need_edit(user)
    from . import gpu_power, power_limit
    from .routers.cluster import _invalidate_power_cache, _maxlps_inputs

    if budget_kw is None and max_rack_kw is None and gpu_power_percent is None:
        return {"ok": False, "error": "say a budget_kw, max_rack_kw, or gpu_power_percent"}
    if budget_kw is not None and float(budget_kw) < 100:
        return {"ok": False, "error": "cluster budget floor is 100 kW. per-rack cap floor is 20 kW."}
    if max_rack_kw is not None and float(max_rack_kw) < 20:
        return {"ok": False, "error": "per-rack cap has to be at least 20 kW"}
    if gpu_power_percent is not None and not (10 <= float(gpu_power_percent) <= 100):
        return {"ok": False, "error": "gpu_power_percent is 10 to 100"}
    power_limit.apply_patch(
        budget_kw=float(budget_kw) if budget_kw is not None else None,
        max_rack_kw=float(max_rack_kw) if max_rack_kw is not None else None,
    )
    gpu_power.apply_loop(
        gpu_power_percent=float(gpu_power_percent) if gpu_power_percent is not None else None,
    )
    if gpu_power_percent is not None:
        gpu_power.LOOP.force = True
        gpu_power.LOOP.last_caps = None
    _invalidate_power_cache()
    snap, _jobs = _maxlps_inputs(db)
    bits = []
    if budget_kw is not None:
        bits.append(f"budget {float(budget_kw):g} kW")
    if max_rack_kw is not None:
        bits.append(f"rack cap {float(max_rack_kw):g} kW")
    if gpu_power_percent is not None:
        bits.append(f"GPU {float(gpu_power_percent):g}%")
    return {
        "ok": True,
        "set": bits,
        "limiter": {
            "budget_kw": (snap or {}).get("budget_kw") if isinstance(snap, dict) else None,
            "max_rack_kw": (snap or {}).get("max_rack_kw") if isinstance(snap, dict) else None,
        },
        "ui": _point_ui("edit", "limiter", "dynamic power", path="/maxlps"),
    }


def tool_add_racks(db: Session, user: User, count: int, prefix: str | None = None, hall: str | None = None, cluster_name: str | None = None, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from . import maxlps_inventory as inv

    created = []
    errors = []
    if cluster_name:
        from .routers.cluster import set_cluster_name

        set_cluster_name(db, cluster_name)
    pre = (prefix or "GB300").strip() or "GB300"
    existing = {r.label.lower() for r in inv.load_racks(db)}
    n = max(1, min(int(count), 18))
    i = 1
    while len(created) < n and i <= n + 40:
        label = f"{pre}-{i:02d}"
        i += 1
        if label.lower() in existing:
            continue
        try:
            rack = inv.create_rack(db, label=label, hall=hall, origin="manual")
            existing.add(label.lower())
            created.append({"id": rack.id, "label": rack.label})
        except Exception as exc:
            errors.append({"label": label, "error": str(exc)})
    labels = [c["label"] for c in created]
    n = len(created)
    return {
        "created": created,
        "errors": errors,
        "cluster_name": cluster_name,
        "ui": _point_ui(
            "add",
            "racks",
            f"adding {n} rack{'s' if n != 1 else ''}",
            path="/cluster?tab=racks",
            ids=labels,
            refresh="before",
        ),
    }


def tool_remove_rack(db: Session, user: User, label: str | None = None, id: str | None = None, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from . import maxlps_inventory as inv

    key = (id or label or "").strip()
    if not key:
        return {"ok": False, "error": "need a rack label or id"}
    out = inv.remove_operator_rack(db, key)
    if out is None:
        return {"ok": False, "error": "rack not found"}
    lab = str(out.get("label") or key)
    return {
        "ok": True,
        "id": out.get("id"),
        "label": lab,
        "nodes": out.get("nodes"),
        "hidden": out.get("hidden"),
        "ui": _point_ui(
            "remove",
            lab,
            f"yeeting {lab}",
            path="/provision",
            ids=[lab, "racks"],
            refresh="after",
        ),
    }


def _match_test(name: str) -> dict[str, Any] | None:
    needle = (name or "").strip().lower()
    if not needle:
        return None
    for t in TEST_CATALOG:
        if t["name"].lower() == needle or t["label"].lower() == needle:
            return t
    for t in TEST_CATALOG:
        if needle in t["name"].lower() or needle in t["label"].lower():
            return t
    return None


def tool_launch_test(db: Session, user: User, name: str, script: str | None = None, id: str | None = None, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from datetime import datetime, timezone

    now = datetime.now(timezone.utc)
    work = db.get(Workload, id) if id else None
    spec = _match_test(name) if not work else None
    if work is None:
        if spec is None:
            return {"ok": False, "error": f"unknown test {name}", "catalog": [t["name"] for t in TEST_CATALOG]}
        copy = (script or "").strip() or DEFAULT_SCRIPTS.get(spec["name"], "")
        work = Workload(
            name=spec["name"],
            kind=spec["kind"],
            status="running",
            project="firmus",
            department="default",
            node_pool="gb300",
            pods_running=spec["pods_requested"],
            pods_requested=spec["pods_requested"],
            gpu_request=spec["gpu_request"],
            gpu_allocation=spec["gpu_request"],
            gpu_mem_gb_request=spec["gpu_mem_gb_request"],
            gpu_mem_gb_alloc=spec["gpu_mem_gb_request"],
            started_at=now,
            detail=spec["detail"] + (" · instance script" if copy else ""),
            script_text=copy or None,
        )
        db.add(work)
    else:
        if script and script.strip():
            work.script_text = script
        elif not work.script_text:
            work.script_text = DEFAULT_SCRIPTS.get(work.name, "") or None
        work.status = "running"
        work.started_at = now
        work.finished_at = None
        work.pods_running = work.pods_requested
        work.gpu_allocation = work.gpu_request
    db.commit()
    db.refresh(work)
    return {
        "ok": True,
        "id": work.id,
        "name": work.name,
        "kind": work.kind,
        "status": work.status,
        "instance_script": bool(work.script_text),
    }


def tool_stop_test(db: Session, user: User, id: str | None = None, name: str | None = None, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from datetime import datetime, timezone
    from sqlalchemy import select

    work = db.get(Workload, id) if id else None
    if work is None and name:
        needle = name.strip().lower()
        rows = list(db.scalars(select(Workload).order_by(Workload.created_at.desc())).all())
        work = next((w for w in rows if w.name.lower() == needle or w.id.lower() == needle), None)
        if work is None:
            work = next((w for w in rows if w.status == "running" and needle in w.name.lower()), None)
    if work is None:
        return {"ok": False, "error": "test run not found"}
    work.status = "stopped"
    work.finished_at = datetime.now(timezone.utc)
    work.detail = (work.detail or "") + " · stopped by operator chat"
    db.commit()
    return {"ok": True, "id": work.id, "name": work.name, "status": work.status}


def tool_sync_inventory(db: Session, user: User, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from . import maxlps_inventory as inv

    try:
        result = inv.sync_from_controller(db)
    except Exception as exc:
        return {"ok": False, "error": str(exc)}
    return {
        "ok": True,
        **result,
        "ui": _point_ui("add", "racks", "syncing Argus", path="/cluster?tab=racks", refresh="before"),
    }


def tool_list_sops(db: Session, **_kw: Any) -> dict[str, Any]:
    from sqlalchemy import select

    from .models import AgentDoc

    rows = list(db.scalars(select(AgentDoc).order_by(AgentDoc.created_at.desc())).all())
    return {
        "docs": [
            {"id": d.id, "title": d.title, "filename": d.filename, "chars": len(d.text or "")}
            for d in rows
        ]
    }


DEFAULT_SCRIPTS: dict[str, str] = {
    "mlperf-gpt3-175b": "#!/bin/bash\nset -euo pipefail\n# instance copy. catalog stays untouched.\nmlperf --model gpt3-175b --config default\n",
    "mlperf-llama2-70b": "#!/bin/bash\nset -euo pipefail\nmlperf --model llama2-70b --config default\n",
    "mlperf-llama3-70b": "#!/bin/bash\nset -euo pipefail\nmlperf --model llama3-70b --config default\n",
    "mlperf-mixtral-8x7b": "#!/bin/bash\nset -euo pipefail\nmlperf --model mixtral-8x7b --config default\n",
    "mlperf-bert-large": "#!/bin/bash\nset -euo pipefail\nmlperf --model bert-large --config default\n",
    "mlperf-resnet50": "#!/bin/bash\nset -euo pipefail\nmlperf --model resnet50 --config default\n",
    "mlperf-retinanet": "#!/bin/bash\nset -euo pipefail\nmlperf --model retinanet --config default\n",
    "mlperf-dlrm-dcnv2": "#!/bin/bash\nset -euo pipefail\nmlperf --model dlrm-dcnv2 --config default\n",
    "mlperf-rnnt": "#!/bin/bash\nset -euo pipefail\nmlperf --model rnnt --config default\n",
    "ib_write_bw": "#!/bin/bash\nset -euo pipefail\nib_write_bw -d mlx5_0 -s 1048576 -n 5000\n",
    "ib_read_bw": "#!/bin/bash\nset -euo pipefail\nib_read_bw -d mlx5_0 -s 1048576 -n 5000\n",
    "ib_send_bw": "#!/bin/bash\nset -euo pipefail\nib_send_bw -d mlx5_0 -s 1048576 -n 5000\n",
    "dcgm-diag": "#!/bin/bash\nset -euo pipefail\ndcgm-diag -r 3\n",
    "nccl-all_reduce": "#!/bin/bash\nset -euo pipefail\nmpirun -np 8 nccl-tests -b 8 -e 8G -f 2 -g 1 all_reduce_perf\n",
    "nccl-all_gather": "#!/bin/bash\nset -euo pipefail\nmpirun -np 8 nccl-tests -b 8 -e 8G -f 2 -g 1 all_gather_perf\n",
}

_RESET_TYPES = {"On", "ForceOff", "GracefulShutdown", "ForceRestart", "GracefulRestart"}
_REDFISH_OK = ("/redfish/v1/", "/redfish/v1")


_IP_RE = r"(?:\d{1,3}\.){3}\d{1,3}"
_MAC_RE = r"(?:[0-9a-fA-F]{2}[:\-]){5}[0-9a-fA-F]{2}|[0-9a-fA-F]{12}"


def parse_node_creds(text: str) -> dict[str, str]:
    """Pull BMC/OS IP and login out of a pasted add-node line."""
    raw = text or ""
    out: dict[str, str] = {}
    os_m = re.search(rf"OS\s*:?\s*({_IP_RE})\s+(\S+)\s*/\s*(\S+)", raw, re.I)
    if os_m:
        out["os_ip"] = os_m.group(1)
        out["os_username"] = os_m.group(2)
        out["os_password"] = os_m.group(3)
    bmc_m = re.search(rf"BMC\s*:?\s*({_IP_RE})\s*:?\s+(\S+)\s*/\s*(\S+)", raw, re.I)
    if bmc_m:
        out["bmc_ip"] = bmc_m.group(1)
        out["bmc_user"] = bmc_m.group(2)
        out["bmc_password"] = bmc_m.group(3)
    mac_m = re.search(rf"(?:BMC\s+)?MAC\s*:?\s*({_MAC_RE})", raw, re.I)
    if mac_m:
        out["bmc_mac"] = mac_m.group(1)
    return out


def _pasted_add_node(text: str) -> dict[str, str] | None:
    parsed = parse_node_creds(text)
    if not (parsed.get("bmc_ip") or parsed.get("bmc_mac")):
        return None
    low = (text or "").lower()
    if re.search(r"\b(add|register|provision)\b", low) or "new node" in low:
        return parsed
    return None


def _inventory_ips(db: Session) -> set[str]:
    from sqlalchemy import select

    from . import maxlps_inventory as inv
    from .models import InventoryNode

    ips: set[str] = set()
    probe = (getattr(settings, "bmc_probe_ip", None) or "").strip()
    if probe:
        ips.add(probe)
    for rack in inv.load_racks(db):
        for n in rack.nodes:
            for raw in (n.bmc_ip, n.os_ip):
                ip = (raw or "").strip()
                if ip:
                    ips.add(ip)
    for n in db.scalars(select(InventoryNode)).all():
        for raw in (n.bmc_ip, n.os_ip):
            ip = (raw or "").strip()
            if ip:
                ips.add(ip)
    return ips


def _inventory_row(node: Any) -> dict[str, Any]:
    return {
        "id": node.id,
        "rack": node.hall_name or "Unassigned",
        "index": 0,
        "hostname": node.name,
        "bmc_ip": (node.bmc_ip or "").strip(),
        "bmc_mac": node.bmc_mac,
        "bmc_user": node.bmc_user or "ADMIN",
        "bmc_password": node.bmc_password or "",
        "os_ip": (node.os_ip or "").strip(),
        "os_username": node.os_username or "root",
        "os_password": node.os_password or "",
        "serial": node.serial,
        "source": getattr(node, "source", None) or "manual",
    }


def _find_inventory_node(db: Session, node: str | None = None) -> Any | None:
    from sqlalchemy import select

    from .models import InventoryNode

    needle = (node or "").strip().lower()
    if not needle:
        return None
    rows = list(db.scalars(select(InventoryNode)).all())
    for n in rows:
        if (n.id or "").lower() == needle:
            return n
    for n in rows:
        blob = [
            n.name,
            n.serial,
            n.bmc_ip,
            n.os_ip,
            n.bmc_mac,
        ]
        if needle in {(str(x or "").strip().lower()) for x in blob if x}:
            return n
    return None


def _find_compute(db: Session, node: str | None = None) -> dict[str, Any] | None:
    from . import maxlps_inventory as inv
    from sqlalchemy import select

    from .models import InventoryNode

    needle = (node or "").strip().lower()
    probe = (getattr(settings, "bmc_probe_ip", None) or "").strip()
    rows: list[dict[str, Any]] = []
    for rack in inv.load_racks(db):
        for n in rack.nodes:
            rows.append(
                {
                    "id": n.id,
                    "rack": rack.label,
                    "index": n.index,
                    "hostname": n.hostname,
                    "bmc_ip": (n.bmc_ip or "").strip(),
                    "bmc_user": n.bmc_user or "root",
                    "bmc_password": n.bmc_password or "",
                    "os_ip": (n.os_ip or "").strip(),
                    "os_username": n.os_username or "root",
                    "os_password": n.os_password or "",
                }
            )
    inv_nodes = list(db.scalars(select(InventoryNode)).all())
    by_bmc = {(x.bmc_ip or "").strip(): x for x in inv_nodes if (x.bmc_ip or "").strip()}
    by_os = {(x.os_ip or "").strip(): x for x in inv_nodes if (x.os_ip or "").strip()}
    for row in rows:
        fallback = by_bmc.get(row["bmc_ip"]) or by_os.get(row["os_ip"])
        if fallback:
            row["os_ip"] = row["os_ip"] or (fallback.os_ip or "").strip()
            row["os_username"] = row["os_username"] if row.get("os_password") else (fallback.os_username or row["os_username"] or "root")
            if not row["os_password"]:
                row["os_password"] = fallback.os_password or ""
            if not row["bmc_password"]:
                row["bmc_password"] = fallback.bmc_password or ""
                row["bmc_user"] = row["bmc_user"] or fallback.bmc_user or "root"
        if row["bmc_ip"] == probe and not row["bmc_password"]:
            row["bmc_user"] = getattr(settings, "bmc_probe_user", None) or row["bmc_user"] or "root"
            row["bmc_password"] = getattr(settings, "bmc_probe_password", None) or ""
    if not needle:
        return next((r for r in rows if r["bmc_ip"] or r["os_ip"]), None)
    for row in rows:
        blob = " ".join(
            str(row.get(k) or "")
            for k in ("id", "rack", "hostname", "bmc_ip", "os_ip")
        ).lower()
        if needle in blob or needle == str(row.get("index")):
            return row
    inv_hit = _find_inventory_node(db, node)
    if inv_hit is not None:
        return _inventory_row(inv_hit)
    if needle == probe.lower():
        return next((r for r in rows if r["bmc_ip"] == probe), None) or {
            "id": "probe",
            "rack": "probe",
            "index": 0,
            "hostname": None,
            "bmc_ip": probe,
            "bmc_user": getattr(settings, "bmc_probe_user", None) or "root",
            "bmc_password": getattr(settings, "bmc_probe_password", None) or "",
            "os_ip": "",
            "os_username": "root",
            "os_password": "",
        }
    return None


def _public_node(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": row.get("id"),
        "rack": row.get("rack"),
        "index": row.get("index"),
        "hostname": row.get("hostname"),
        "bmc_ip": row.get("bmc_ip") or None,
        "bmc_user": row.get("bmc_user"),
        "bmc_login": bool(row.get("bmc_password")),
        "os_ip": row.get("os_ip") or None,
        "os_user": row.get("os_username") or "root",
        "os_login": bool(row.get("os_password")),
    }


def tool_list_nodes(db: Session, **_kw: Any) -> dict[str, Any]:
    from sqlalchemy import select

    from . import maxlps_inventory as inv
    from .models import InventoryNode

    nodes = []
    seen: set[str] = set()
    for rack in inv.load_racks(db):
        for n in rack.nodes:
            if not ((n.bmc_ip or "").strip() or (n.os_ip or "").strip() or (n.hostname or "").strip()):
                continue
            row = _find_compute(db, (n.bmc_ip or n.os_ip or n.id) or "")
            if row:
                key = str(row.get("bmc_ip") or row.get("os_ip") or row.get("id") or "")
                if key:
                    seen.add(key)
                nodes.append(_public_node(row))
    for n in db.scalars(select(InventoryNode)).all():
        if not ((n.bmc_ip or "").strip() or (n.os_ip or "").strip()):
            continue
        key = (n.bmc_ip or "").strip() or (n.os_ip or "").strip() or n.id
        if key in seen:
            continue
        seen.add(key)
        nodes.append(_public_node(_inventory_row(n)))
    if not nodes:
        probe = _find_compute(db, getattr(settings, "bmc_probe_ip", None))
        if probe:
            nodes.append(_public_node(probe))
    return {"nodes": nodes}


def tool_bmc_get(db: Session, node: str, path: str, **_kw: Any) -> dict[str, Any]:
    from . import bmc_health

    row = _find_compute(db, node)
    if row is None:
        return {"ok": False, "error": "node not in inventory"}
    ip = row.get("bmc_ip") or ""
    if not ip:
        return {"ok": False, "error": "no BMC IP"}
    if ip not in _inventory_ips(db):
        return {"ok": False, "error": "BMC IP is not in inventory"}
    if not row.get("bmc_password"):
        return {"ok": False, "error": "no BMC login stored"}
    p = (path or "").strip() or "/redfish/v1"
    if ".." in p or not p.startswith("/redfish/v1"):
        return {"ok": False, "error": "path must be under /redfish/v1"}
    body = bmc_health.redfish_get(ip, p, row["bmc_user"], row["bmc_password"])
    err = body.get("_error") if isinstance(body, dict) else None
    slim = {k: v for k, v in (body or {}).items() if not str(k).startswith("_")} if isinstance(body, dict) else {}
    if err:
        return {"ok": False, "bmc_ip": ip, "path": p, "error": err}
    return {"ok": True, "bmc_ip": ip, "path": p, "via": "direct", "body": _trim_redfish(slim)}


def _trim_redfish(body: dict[str, Any], limit: int = 3500) -> dict[str, Any]:
    text = json.dumps(body, default=str)
    if len(text) <= limit:
        return body
    keep = {k: body[k] for k in list(body)[:24]}
    keep["_truncated"] = True
    return keep


def tool_bmc_reset(db: Session, user: User, node: str, reset_type: str, **_kw: Any) -> dict[str, Any]:
    _need_edit(user)
    from . import bmc_health

    kind = (reset_type or "").strip()
    if kind not in _RESET_TYPES:
        return {"ok": False, "error": f"reset_type must be one of {sorted(_RESET_TYPES)}"}
    row = _find_compute(db, node)
    if row is None:
        return {"ok": False, "error": "node not in inventory"}
    ip = row.get("bmc_ip") or ""
    if not ip or ip not in _inventory_ips(db):
        return {"ok": False, "error": "BMC IP is not in inventory"}
    if not row.get("bmc_password"):
        return {"ok": False, "error": "no BMC login stored"}
    path = "/redfish/v1/Systems/System_0/Actions/ComputerSystem.Reset"
    body = bmc_health.redfish_call(
        ip,
        path,
        row["bmc_user"],
        row["bmc_password"],
        method="POST",
        body={"ResetType": kind},
    )
    err = body.get("_error") if isinstance(body, dict) else None
    if err:
        return {"ok": False, "bmc_ip": ip, "reset_type": kind, "error": err, "via": "direct"}
    return {"ok": True, "bmc_ip": ip, "reset_type": kind, "status": body.get("_status"), "via": "direct"}


def _console_ui(
    row: dict[str, Any],
    command: str = "",
    stdout: str = "",
    stderr: str = "",
    exit_code: int | None = None,
    error: str | None = None,
    ok: bool = True,
    host: str | None = None,
    user: str | None = None,
) -> dict[str, Any]:
    tag = str(row.get("rack") or row.get("hostname") or row.get("os_ip") or row.get("bmc_ip") or "node")
    return {
        "kind": "console",
        "node": tag,
        "host": host if host is not None else (row.get("os_ip") or ""),
        "user": user if user is not None else (row.get("os_username") or "root"),
        "command": command,
        "stdout": stdout or "",
        "stderr": stderr or "",
        "exit": exit_code,
        "error": error,
        "ok": ok,
    }


def _ssh_login(row: dict[str, Any], via: str | None = None) -> dict[str, str]:
    want = (via or "").strip().lower()
    bmc = {
        "host": (row.get("bmc_ip") or "").strip(),
        "user": (row.get("bmc_user") or "root"),
        "password": row.get("bmc_password") or "",
        "via": "bmc",
    }
    os_login = {
        "host": (row.get("os_ip") or "").strip(),
        "user": (row.get("os_username") or "root"),
        "password": row.get("os_password") or "",
        "via": "os",
    }
    order = [bmc, os_login] if want == "bmc" else [os_login, bmc]
    for cand in order:
        if cand["host"] and cand["password"]:
            return cand
    empty = bmc if want == "bmc" else os_login
    return empty


def tool_open_console(
    db: Session, user: User, node: str, command: str | None = None, via: str | None = None, **_kw: Any
) -> dict[str, Any]:
    row = _find_compute(db, node)
    if row is None:
        return {"ok": False, "error": "node not in inventory"}
    cmd = (command or "").strip()
    if cmd:
        return tool_os_exec(db, user, node=node, command=cmd, via=via)
    login = _ssh_login(row, via)
    err = None
    if not login["host"]:
        err = "no BMC or OS IP stored"
    elif not login["password"]:
        err = f"no {login['via']} password stored"
    elif login["host"] not in _inventory_ips(db):
        err = "IP is not in inventory"
    ui = _console_ui(row, error=err, ok=not err, host=login["host"], user=login["user"])
    return {"ok": not err, "host": ui["host"], "user": ui["user"], "via": login["via"], "error": err, "ui": ui}


def tool_os_exec(
    db: Session, user: User, node: str, command: str, via: str | None = None, **_kw: Any
) -> dict[str, Any]:
    _need_edit(user)
    import paramiko

    cmd = (command or "").strip()
    row = _find_compute(db, node)
    if row is None:
        return {"ok": False, "error": "node not in inventory"}
    login = _ssh_login(row, via)
    if not cmd:
        ui = _console_ui(row, error="empty command", ok=False, host=login["host"], user=login["user"])
        return {"ok": False, "error": "empty command", "ui": ui}
    if len(cmd) > 2000:
        ui = _console_ui(row, command=cmd[:80], error="command too long", ok=False, host=login["host"], user=login["user"])
        return {"ok": False, "error": "command too long", "ui": ui}
    host = login["host"]
    if not host:
        err = "no BMC or OS IP stored"
        return {"ok": False, "error": err, "ui": _console_ui(row, command=cmd, error=err, ok=False)}
    if host not in _inventory_ips(db):
        err = "IP is not in inventory"
        return {"ok": False, "error": err, "ui": _console_ui(row, command=cmd, error=err, ok=False, host=host, user=login["user"])}
    if not login["password"]:
        err = f"no {login['via']} password stored"
        return {"ok": False, "error": err, "ui": _console_ui(row, command=cmd, error=err, ok=False, host=host, user=login["user"])}

    def _exec(target: str, port: int) -> tuple[str, str, int]:
        client = paramiko.SSHClient()
        client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        try:
            client.connect(
                target,
                port=port,
                username=login["user"],
                password=login["password"],
                timeout=12,
                allow_agent=False,
                look_for_keys=False,
                banner_timeout=12,
                auth_timeout=12,
            )
            _stdin, stdout, stderr = client.exec_command(cmd, timeout=60)
            return (
                stdout.read().decode("utf-8", "replace"),
                stderr.read().decode("utf-8", "replace"),
                stdout.channel.recv_exit_status(),
            )
        finally:
            try:
                client.close()
            except Exception:
                pass

    hopped = False
    try:
        out, err_text, code = _exec(host, 22)
    except Exception as exc:
        err = str(exc)[:180]
        reachable = any(bit in err.lower() for bit in ("timed out", "timeout", "unreachable", "no route", "refused", "reset"))
        if login["via"] != "os" or not reachable:
            if login["via"] == "bmc":
                err = f"BMC SSH failed ({err}). BMC is Redfish over HTTPS. Use bmc_get / bmc_reset, or SSH the OS IP."
            return {
                "ok": False,
                "host": host,
                "via": login["via"],
                "error": err,
                "ui": _console_ui(row, command=cmd, error=err, ok=False, host=host, user=login["user"]),
            }
        try:
            from . import pxe_hop

            from .pxe_servers import pxe_for_address

            pxe = pxe_for_address(host)
            if pxe and not pxe_hop.ensure_host(pxe):
                err = f"OS SSH timed out from this PC. Could not hop through {pxe}."
                return {
                    "ok": False,
                    "host": host,
                    "via": "os",
                    "error": err,
                    "ui": _console_ui(row, command=cmd, error=err, ok=False, host=host, user=login["user"]),
                }
            if not pxe and not pxe_hop.status().get("connected"):
                err = "OS SSH timed out from this PC. Connect the PXE hop, then run the command again."
                return {
                    "ok": False,
                    "host": host,
                    "via": "os",
                    "error": err,
                    "ui": _console_ui(row, command=cmd, error=err, ok=False, host=host, user=login["user"]),
                }
            if pxe:
                with pxe_hop.on_host(pxe):
                    local = pxe_hop.open_https_tunnel(f"ssh-{host}", host, 22)
            else:
                local = pxe_hop.open_https_tunnel(f"ssh-{host}", host, 22)
            out, err_text, code = _exec("127.0.0.1", local)
            hopped = True
        except Exception as hop_exc:
            err = f"OS SSH through the PXE hop failed ({hop_exc})"[:180]
            return {
                "ok": False,
                "host": host,
                "via": "hop",
                "error": err,
                "ui": _console_ui(row, command=cmd, error=err, ok=False, host=host, user=login["user"]),
            }
    ok = code == 0
    ui = _console_ui(
        row, command=cmd, stdout=out[-6000:], stderr=err_text[-2000:], exit_code=code, ok=ok, host=host, user=login["user"]
    )
    return {
        "ok": ok,
        "host": host,
        "user": login["user"],
        "via": "hop" if hopped else login["via"],
        "exit": code,
        "stdout": out[-6000:],
        "stderr": err_text[-2000:],
        "ui": ui,
    }


def tool_get_test_script(db: Session, name: str | None = None, id: str | None = None, **_kw: Any) -> dict[str, Any]:
    work = db.get(Workload, id) if id else None
    if work is not None:
        script = work.script_text or DEFAULT_SCRIPTS.get(work.name, "")
        return {
            "ok": True,
            "id": work.id,
            "name": work.name,
            "instance": bool(work.script_text),
            "script": script,
        }
    spec = _match_test(name or "")
    if spec is None:
        return {"ok": False, "error": "unknown test", "catalog": [t["name"] for t in TEST_CATALOG]}
    return {
        "ok": True,
        "name": spec["name"],
        "instance": False,
        "script": DEFAULT_SCRIPTS.get(spec["name"], ""),
    }


def tool_save_test_script(
    db: Session,
    user: User,
    script: str,
    name: str | None = None,
    id: str | None = None,
    **_kw: Any,
) -> dict[str, Any]:
    _need_edit(user)

    text = (script or "").replace("\x00", "")
    if not text.strip():
        return {"ok": False, "error": "empty script"}
    if len(text) > 80000:
        text = text[:80000]
    work = db.get(Workload, id) if id else None
    if work is None:
        spec = _match_test(name or "")
        if spec is None:
            return {"ok": False, "error": "give a catalog name or run id"}
        work = Workload(
            name=spec["name"],
            kind=spec["kind"],
            status="pending",
            project="firmus",
            department="default",
            node_pool="gb300",
            pods_running=0,
            pods_requested=spec["pods_requested"],
            gpu_request=spec["gpu_request"],
            gpu_allocation=0,
            gpu_mem_gb_request=spec["gpu_mem_gb_request"],
            gpu_mem_gb_alloc=0,
            started_at=None,
            detail=f"{spec['detail']} · instance script only",
            script_text=text,
        )
        db.add(work)
    else:
        work.script_text = text
        extra = " · instance script only"
        if extra not in (work.detail or ""):
            work.detail = (work.detail or "") + extra
    db.commit()
    db.refresh(work)
    return {"ok": True, "id": work.id, "name": work.name, "status": work.status, "instance_only": True, "chars": len(text)}


_ALLOWED_PATHS = {
    "/cluster",
    "/maxlps",
    "/network",
    "/provision",
    "/testing",
    "/settings",
    "/users",
    "/dashboard",
    "/calendar",
    "/tasks",
    "/task-calendar",
}
_PAGE_ALIASES: dict[str, tuple[str, str | None]] = {
    "cluster": ("/cluster", None),
    "overview": ("/cluster", "Overview"),
    "racks": ("/cluster", "Racks"),
    "floor": ("/cluster", "Floor"),
    "maxlps": ("/maxlps", None),
    "dynamic power": ("/maxlps", None),
    "power": ("/maxlps", None),
    "limiter": ("/maxlps", None),
    "limitter": ("/maxlps", None),
    "power limiter": ("/maxlps", None),
    "power limitter": ("/maxlps", None),
    "power limit": ("/maxlps", None),
    "power cap": ("/maxlps", None),
    "lps": ("/maxlps", None),
    "network": ("/network", None),
    "provision": ("/provision", None),
    "testing": ("/testing", None),
    "tests": ("/testing", None),
    "sanity": ("/testing", None),
    "bom": ("/testing", None),
    "golden": ("/testing", None),
    "reports": ("/testing", None),
    "settings": ("/settings", None),
    "users": ("/users", None),
    "profile": ("/users", None),
    "dashboard": ("/dashboard", None),
    "calendar": ("/calendar", None),
    "tasks": ("/tasks", None),
    "task-calendar": ("/task-calendar", None),
}
_POWER_PAGE_HINTS = (
    "maxlps",
    "max lps",
    "dynamic power",
    "limiter",
    "limitter",
    "power limiter",
    "power limitter",
    "power limit",
    "power cap",
    "power budget",
    "gpu power",
    "closed loop",
    "closedloop",
)
_BMC_POWER_HINTS = (
    "power on",
    "power off",
    "power cycle",
    "force off",
    "forceoff",
    "graceful shutdown",
    "graceful restart",
)


def _fold_page(raw: str) -> str:
    return " ".join((raw or "").lower().replace("_", " ").replace("-", " ").replace("/", " ").split())


def _wants_dynamic_power(raw: str) -> bool:
    t = _fold_page(raw)
    if not t:
        return False
    if any(b in t for b in _BMC_POWER_HINTS):
        return False
    if t in {"power", "lps", "watts", "watt"}:
        return True
    if any(h in t for h in _POWER_PAGE_HINTS):
        return True
    return "power" in t.split()


def _norm_page(raw: str) -> tuple[str, str | None]:
    s = (raw or "").strip()
    if not s:
        return "", None
    low = s.lower().strip()
    key = low.strip("/")
    folded = _fold_page(s)
    if any(b in folded for b in _BMC_POWER_HINTS):
        return "", None
    if _wants_dynamic_power(s):
        return "/maxlps", None
    if key in _PAGE_ALIASES:
        return _PAGE_ALIASES[key]
    if low.startswith("/"):
        path = "/" + low.strip("/").split("?")[0].split("#")[0]
        if path in _ALLOWED_PATHS:
            tab = None
            if "tab=racks" in low or path == "/cluster" and "racks" in low:
                tab = "Racks"
            elif "tab=floor" in low or "floor" in low and path == "/cluster":
                tab = "Floor"
            elif "tab=overview" in low:
                tab = "Overview"
            return path, tab
        return "", None
    for alias, dest in _PAGE_ALIASES.items():
        if alias in key or key in alias:
            return dest
    return "", None


def _point_ui(
    action: str,
    target: str,
    label: str,
    *,
    path: str | None = None,
    ids: list[str] | None = None,
    refresh: str | None = None,
) -> list[dict[str, Any]]:
    ui: list[dict[str, Any]] = []
    if path:
        dest, tab = _norm_page(path)
        if dest:
            href = dest
            if tab == "Racks":
                href = "/cluster?tab=racks"
            elif tab == "Floor":
                href = "/cluster?tab=floor"
            elif tab == "Overview":
                href = "/cluster?tab=overview"
            ui.append({"kind": "navigate", "path": href})
            if tab:
                ui.append({"kind": "click", "target": tab})
    cmd: dict[str, Any] = {
        "kind": "highlight",
        "action": action,
        "target": (target or "")[:120],
        "label": (label or "")[:80],
        "ids": [str(x)[:80] for x in (ids or [])][:18],
    }
    if refresh:
        cmd["refresh"] = refresh
    ui.append(cmd)
    return ui


def tool_open_page(path: str, **_kw: Any) -> dict[str, Any]:
    dest, tab = _norm_page(path)
    if not dest:
        return {"ok": False, "error": "unknown page", "pages": sorted(_ALLOWED_PATHS)}
    href = dest
    if tab == "Racks":
        href = "/cluster?tab=racks"
    elif tab == "Floor":
        href = "/cluster?tab=floor"
    elif tab == "Overview":
        href = "/cluster?tab=overview"
    ui: list[dict[str, Any]] = [{"kind": "navigate", "path": href}]
    if tab:
        ui.append({"kind": "click", "target": tab})
    return {"ok": True, "path": href, "tab": tab, "ui": ui}


def tool_ui_fill(target: str, value: str, submit: bool = False, **_kw: Any) -> dict[str, Any]:
    needle = (target or "").strip()
    if not needle:
        return {"ok": False, "error": "empty target"}
    return {
        "ok": True,
        "ui": {"kind": "fill", "target": needle[:120], "value": str(value)[:4000], "submit": bool(submit)},
    }


def tool_ui_click(target: str, **_kw: Any) -> dict[str, Any]:
    needle = (target or "").strip()
    if not needle:
        return {"ok": False, "error": "empty target"}
    return {"ok": True, "ui": {"kind": "click", "target": needle[:120], "label": needle[:80]}}


_HIGHLIGHT_PAGES = {
    "health": "/cluster?tab=overview",
    "overview": "/cluster?tab=overview",
    "racks": "/cluster?tab=racks",
    "rack": "/cluster?tab=racks",
    "floor": "/cluster?tab=floor",
    "testing": "/testing",
    "tests": "/testing",
    "sanity": "/testing",
    "bom": "/testing",
    "golden": "/testing",
    "reports": "/testing",
    "limiter": "/maxlps",
    "limitter": "/maxlps",
    "maxlps": "/maxlps",
    "power": "/maxlps",
    "lps": "/maxlps",
    "dynamic power": "/maxlps",
    "power limiter": "/maxlps",
    "power limitter": "/maxlps",
    "power cap": "/maxlps",
    "network": "/network",
    "provision": "/provision",
    "settings": "/settings",
    "users": "/users",
    "cluster-name": "/cluster",
    "register": "/cluster?tab=racks",
}


def tool_highlight_ui(target: str, label: str | None = None, action: str | None = None, **_kw: Any) -> dict[str, Any]:
    needle = (target or "").strip()
    if not needle:
        return {"ok": False, "error": "empty target"}
    act = (action or "look").strip().lower()
    if act not in {"look", "add", "remove", "edit", "click"}:
        act = "look"
    low = needle.lower()
    page = _HIGHLIGHT_PAGES.get(low)
    if page is None and _wants_dynamic_power(needle):
        page = "/maxlps"
    point_at = {
        "health": "health",
        "overview": "overview",
        "rack": "racks",
        "tests": "catalog",
        "testing": "catalog",
        "catalog": "catalog",
        "sanity": "sanity",
        "bom": "sanity",
        "golden": "golden",
        "reports": "reports",
        "power": "limiter",
        "maxlps": "limiter",
        "lps": "limiter",
        "limiter": "limiter",
        "limitter": "limiter",
        "dynamic power": "limiter",
        "power limiter": "limiter",
        "power limitter": "limiter",
        "power cap": "limiter",
    }.get(low, needle)
    if page == "/maxlps":
        point_at = "limiter"
    caption = (label or "").strip() or f"this is {needle}"
    return {"ok": True, "ui": _point_ui(act, point_at, caption[:80], path=page, ids=[needle])}


def tool_patch_rack(
    db: Session,
    user: User,
    rack: str,
    label: str | None = None,
    hall: str | None = None,
    serial: str | None = None,
    notes: str | None = None,
    **_kw: Any,
) -> dict[str, Any]:
    _need_edit(user)
    from . import maxlps_inventory as inv

    row = None
    if rack:
        needle = rack.strip().lower()
        row = inv.get_rack(db, rack) if len(rack) >= 8 else None
        if row is None:
            row = next((r for r in inv.load_racks(db) if r.label.lower() == needle or r.id.lower() == needle), None)
    if row is None:
        return {"ok": False, "error": "rack not found"}
    body: dict[str, Any] = {}
    if label is not None:
        body["label"] = label
    if hall is not None:
        body["hall"] = hall
    if serial is not None:
        body["serial"] = serial
    if notes is not None:
        body["notes"] = notes
    if not body:
        return {"ok": False, "error": "nothing to patch", "id": row.id, "label": row.label}
    try:
        updated = inv.patch_rack(db, row, body)
    except ValueError as exc:
        return {"ok": False, "error": str(exc)}
    return {
        "ok": True,
        "id": updated.id,
        "label": updated.label,
        "hall": updated.hall,
        "serial": updated.serial,
        "ui": _point_ui("edit", updated.label, f"patched {updated.label}", path="/cluster?tab=racks", ids=[updated.label]),
    }


def _node_fields(
    *,
    bmc_ip: str | None = None,
    bmc_mac: str | None = None,
    bmc_user: str | None = None,
    bmc_password: str | None = None,
    os_ip: str | None = None,
    os_username: str | None = None,
    os_password: str | None = None,
    serial: str | None = None,
    name: str | None = None,
    hostname: str | None = None,
    hall_name: str | None = None,
    rack: str | None = None,
    raw: str | None = None,
) -> dict[str, str]:
    parsed = parse_node_creds(raw or "")
    merged: dict[str, str] = {}
    for key, val in (
        ("bmc_ip", bmc_ip),
        ("bmc_mac", bmc_mac),
        ("bmc_user", bmc_user),
        ("bmc_password", bmc_password),
        ("os_ip", os_ip),
        ("os_username", os_username),
        ("os_password", os_password),
        ("serial", serial),
        ("name", name or hostname),
        ("hall_name", hall_name or rack),
    ):
        pick = val if val is not None and str(val).strip() else parsed.get(key)
        if pick is not None and str(pick).strip():
            merged[key] = str(pick).strip()
    return merged


def _public_added(node: Any, *, created: bool, filled: dict[str, Any] | None = None) -> dict[str, Any]:
    tag = node.name or node.serial or node.bmc_ip or "node"
    ids = [x for x in (node.name, node.serial, node.bmc_ip, node.os_ip, node.id) if x]
    notes = [n for n in ((filled or {}).get("notes") or []) if isinstance(n, str)]
    return {
        "ok": True,
        "created": created,
        "id": node.id,
        "name": node.name,
        "serial": node.serial,
        "bmc_ip": node.bmc_ip,
        "bmc_mac": node.bmc_mac,
        "bmc_user": node.bmc_user,
        "bmc_login": bool(node.bmc_password),
        "os_ip": node.os_ip,
        "os_user": node.os_username,
        "os_login": bool(node.os_password),
        "rack": node.hall_name,
        "status": node.provision_status,
        "notes": notes,
        "ui": _point_ui(
            "add" if created else "edit",
            tag,
            f"{'added' if created else 'patched'} {tag}",
            path="/provision",
            ids=ids,
            refresh="before",
        ),
        "choices": [
            {
                "label": f"Open console on {tag}",
                "tool": "open_console",
                "arguments": {"node": node.os_ip or node.bmc_ip or node.name},
            },
            {"label": "Show connected nodes", "tool": "list_nodes", "arguments": {}},
        ],
    }


def tool_patch_node(
    db: Session,
    user: User,
    node: str,
    bmc_ip: str | None = None,
    bmc_mac: str | None = None,
    bmc_user: str | None = None,
    bmc_password: str | None = None,
    os_ip: str | None = None,
    os_username: str | None = None,
    os_password: str | None = None,
    hostname: str | None = None,
    name: str | None = None,
    serial: str | None = None,
    hall_name: str | None = None,
    rack: str | None = None,
    **_kw: Any,
) -> dict[str, Any]:
    _need_edit(user)
    from . import maxlps_inventory as inv
    from .models import MaxLpsNode
    from .provision import apply_node_fields, commit_enriched

    body = _node_fields(
        bmc_ip=bmc_ip,
        bmc_mac=bmc_mac,
        bmc_user=bmc_user,
        bmc_password=bmc_password,
        os_ip=os_ip,
        os_username=os_username,
        os_password=os_password,
        serial=serial,
        name=name,
        hostname=hostname,
        hall_name=hall_name,
        rack=rack,
    )
    if not body:
        return {"ok": False, "error": "nothing to patch"}
    found = _find_compute(db, node)
    inv_node = _find_inventory_node(db, node)
    if inv_node is None and found:
        inv_node = _find_inventory_node(db, found.get("bmc_ip") or found.get("os_ip") or found.get("id"))
    patched = False
    if found and found.get("id") and found.get("id") != "probe":
        row = db.get(MaxLpsNode, found.get("id"))
        if row is not None:
            max_body = dict(body)
            if "name" in max_body and "hostname" not in max_body:
                max_body["hostname"] = max_body.pop("name")
            if max_body:
                inv.patch_node(db, row, max_body)
                patched = True
    if inv_node is not None:
        apply_node_fields(db, inv_node, **body)
        if "hall_name" in body:
            try:
                inv.ensure_rack_for_node(db, inv_node)
            except Exception:
                pass
        if {"bmc_ip", "bmc_mac", "bmc_user", "bmc_password"} & set(body):
            filled = commit_enriched(db, inv_node)
        else:
            db.commit()
            db.refresh(inv_node)
            filled = {}
        patched = True
        return _public_added(inv_node, created=False, filled=filled)
    if not patched:
        return {"ok": False, "error": "node not in inventory. Use add_node."}
    fresh = _find_compute(db, found.get("id") if found else node) or found or {}
    return {"ok": True, **_public_node(fresh)}


def tool_add_node(
    db: Session,
    user: User,
    bmc_ip: str | None = None,
    bmc_mac: str | None = None,
    bmc_user: str | None = None,
    bmc_password: str | None = None,
    os_ip: str | None = None,
    os_username: str | None = None,
    os_password: str | None = None,
    serial: str | None = None,
    name: str | None = None,
    hall_name: str | None = None,
    rack: str | None = None,
    raw: str | None = None,
    **_kw: Any,
) -> dict[str, Any]:
    _need_edit(user)
    from .provision import add_or_update_node

    fields = _node_fields(
        bmc_ip=bmc_ip,
        bmc_mac=bmc_mac,
        bmc_user=bmc_user,
        bmc_password=bmc_password,
        os_ip=os_ip,
        os_username=os_username,
        os_password=os_password,
        serial=serial,
        name=name,
        hall_name=hall_name,
        rack=rack,
        raw=raw,
    )
    if not fields.get("bmc_ip") and not fields.get("bmc_mac"):
        return {"ok": False, "error": "Need a BMC IP or BMC MAC"}
    try:
        node, filled, created = add_or_update_node(db, **fields)
    except ValueError as exc:
        return {"ok": False, "error": str(exc)}
    return _public_added(node, created=created, filled=filled)


def tool_offer_choices(choices: list[Any] | None = None, **_kw: Any) -> dict[str, Any]:
    allowed = set(HANDLERS) - {"offer_choices"}
    out = []
    for raw in choices or []:
        if not isinstance(raw, dict):
            continue
        label = str(raw.get("label") or "").strip()
        tool = str(raw.get("tool") or "").strip()
        if not label or tool not in allowed:
            continue
        args = raw.get("arguments") if isinstance(raw.get("arguments"), dict) else {}
        hint = str(raw.get("hint") or "").strip()[:140]
        out.append(
            {
                "label": label[:90],
                "tool": tool,
                "arguments": args,
                "hint": hint or None,
                "mutating": tool in MUTATING,
            }
        )
        if len(out) >= 4:
            break
    return {"ok": True, "choices": out}


def tool_read_sop(db: Session, id: str | None = None, title: str | None = None, **_kw: Any) -> dict[str, Any]:
    from sqlalchemy import select

    from .models import AgentDoc

    doc = db.get(AgentDoc, id) if id else None
    if doc is None and title:
        needle = title.strip().lower()
        rows = list(db.scalars(select(AgentDoc)).all())
        doc = next((d for d in rows if d.title.lower() == needle or d.filename.lower() == needle), None)
        if doc is None:
            doc = next((d for d in rows if needle in d.title.lower() or needle in d.filename.lower()), None)
    if doc is None:
        return {"ok": False, "error": "SOP not found", **tool_list_sops(db)}
    text = doc.text or ""
    if len(text) > 12000:
        text = text[:12000] + "\n[truncated]"
    return {"ok": True, "id": doc.id, "title": doc.title, "filename": doc.filename, "text": text}


HANDLERS: dict[str, Callable[..., Any]] = {
    "get_cluster_status": tool_get_cluster_status,
    "get_health": tool_get_health,
    "get_configuration": tool_get_configuration,
    "list_tests": tool_list_tests,
    "list_test_reports": tool_list_test_reports,
    "read_test_logs": tool_read_test_logs,
    "run_argus_test": tool_run_argus_test,
    "start_test_agent": tool_start_test_agent,
    "list_agents": tool_list_agents,
    "stop_agent": tool_stop_agent,
    "open_agent": tool_open_agent,
    "list_sanity": tool_list_sanity,
    "run_sanity_check": tool_run_sanity_check,
    "capture_golden": tool_capture_golden,
    "generate_l11_report": tool_generate_l11_report,
    "list_test_runs": tool_list_test_runs,
    "read_logs": tool_read_logs,
    "set_cluster_name": tool_set_cluster_name,
    "set_power_cap": tool_set_power_cap,
    "add_racks": tool_add_racks,
    "remove_rack": tool_remove_rack,
    "launch_test": tool_launch_test,
    "stop_test": tool_stop_test,
    "sync_inventory": tool_sync_inventory,
    "list_sops": tool_list_sops,
    "read_sop": tool_read_sop,
    "list_nodes": tool_list_nodes,
    "bmc_get": tool_bmc_get,
    "bmc_reset": tool_bmc_reset,
    "open_console": tool_open_console,
    "os_exec": tool_os_exec,
    "get_test_script": tool_get_test_script,
    "save_test_script": tool_save_test_script,
    "open_page": tool_open_page,
    "ui_fill": tool_ui_fill,
    "ui_click": tool_ui_click,
    "highlight_ui": tool_highlight_ui,
    "patch_rack": tool_patch_rack,
    "patch_node": tool_patch_node,
    "add_node": tool_add_node,
    "offer_choices": tool_offer_choices,
    "list_attachments": tool_list_attachments,
    "read_attachment": tool_read_attachment,
    "list_app_routes": tool_list_app_routes,
    "call_app": tool_call_app,
}

MUTATING = {
    "set_cluster_name",
    "set_power_cap",
    "add_racks",
    "remove_rack",
    "launch_test",
    "run_argus_test",
    "start_test_agent",
    "stop_agent",
    "run_sanity_check",
    "capture_golden",
    "generate_l11_report",
    "stop_test",
    "sync_inventory",
    "bmc_reset",
    "open_console",
    "os_exec",
    "save_test_script",
    "patch_rack",
    "patch_node",
    "add_node",
    "ui_fill",
}

_DEFAULT_CHOICES = [
    {"label": "Check live health", "tool": "get_health", "arguments": {}, "hint": "Leak, temp, fans, PSU"},
    {"label": "Show connected nodes", "tool": "list_nodes", "arguments": {}, "hint": "BMC and OS logins"},
    {"label": "List tests I can run", "tool": "list_tests", "arguments": {}, "hint": "MLPerf, IB, NCCL"},
]


def _fallback_choices(db: Session, actions: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    """Tap targets when Grok skips offer_choices. Never includes BMC reset."""
    names = {str(a.get("name") or "") for a in (actions or [])}
    nodes: list[dict[str, Any]] = []
    try:
        nodes = list((tool_list_nodes(db) or {}).get("nodes") or [])
    except Exception:
        nodes = []
    live = next((n for n in nodes if n.get("bmc_login") or n.get("os_login")), None)
    tag = "node"
    node_key = ""
    if live:
        tag = str(live.get("rack") or live.get("hostname") or live.get("bmc_ip") or "node")
        node_key = str(live.get("bmc_ip") or live.get("os_ip") or live.get("rack") or tag)

    picks: list[dict[str, Any]] = []

    def add(label: str, tool: str, arguments: dict[str, Any] | None = None, hint: str | None = None) -> None:
        if len(picks) >= 4:
            return
        args = arguments or {}
        if tool not in HANDLERS or tool == "offer_choices":
            return
        if any(p["label"] == label or (p["tool"] == tool and p.get("arguments") == args) for p in picks):
            return
        picks.append(
            {
                "label": label[:90],
                "tool": tool,
                "arguments": args,
                "hint": (hint or "")[:140] or None,
                "mutating": tool in MUTATING,
            }
        )

    for a in actions or []:
        if str(a.get("name") or "") != "start_test_agent" or not a.get("ok"):
            continue
        blob = a.get("summary") or ""
        row: dict[str, Any] = blob if isinstance(blob, dict) else {}
        if isinstance(blob, str):
            try:
                parsed = json.loads(blob)
            except json.JSONDecodeError:
                parsed = {}
            if isinstance(parsed, dict):
                row = parsed
        minion = str(row.get("minion") or "minion")
        jid = str(row.get("id") or "")
        if jid:
            add(f"See {minion}", "open_agent", {"id": jid}, "minion radio")

    if "get_health" not in names:
        add("Check live health", "get_health", {}, "Leak, temp, fans, PSU")
    if live and live.get("bmc_login"):
        add(
            f"Read BMC on {tag}",
            "bmc_get",
            {"node": node_key, "path": "/redfish/v1/Systems/System_0"},
            "Power, model, serial",
        )
    if live and live.get("os_login"):
        add(
            f"Open SSH on {tag}",
            "open_console",
            {"node": str(live.get("os_ip") or node_key), "command": "uname -a; uptime"},
            "Root OS console",
        )
    elif "list_nodes" not in names:
        add("Show connected nodes", "list_nodes", {}, "BMC and OS logins")
    if "list_tests" in names:
        add("Launch DCGM on a new instance", "launch_test", {"name": "dcgm-diag"}, "Catalog stays untouched")
        if "start_test_agent" not in names:
            add("Send a sanity minion", "start_test_agent", {"kind": "sanity"}, "Crew radio in AGENTS")
        add("Show recent test runs", "list_test_runs", {"limit": 8}, "Instance scripts live here")
    elif "save_test_script" in names:
        add("Show recent test runs", "list_test_runs", {"limit": 8}, "Find the instance copy")
        add("List tests I can run", "list_tests", {}, "MLPerf, IB, NCCL")
    else:
        add("List tests I can run", "list_tests", {}, "MLPerf, IB, NCCL")
    if not picks:
        for row in _DEFAULT_CHOICES:
            add(str(row["label"]), str(row["tool"]), dict(row.get("arguments") or {}), row.get("hint"))
    return picks[:4]


def _xai_post(payload: dict[str, Any], key: str) -> dict[str, Any]:
    if not key:
        raise RuntimeError("Grok key is not set")
    body = json.dumps(payload).encode("utf-8")
    req = urllib.request.Request(
        f"{_XAI}/responses",
        data=body,
        method="POST",
        headers={"Authorization": f"Bearer {key}", "Content-Type": "application/json"},
    )
    try:
        with urllib.request.urlopen(req, context=_TLS, timeout=90) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        err = exc.read().decode("utf-8", errors="replace")[:800]
        raise RuntimeError(f"Grok API {exc.code}: {err}") from exc


def _fn_calls(resp: dict[str, Any]) -> list[dict[str, Any]]:
    out = []
    for item in resp.get("output") or []:
        if not isinstance(item, dict):
            continue
        kind = str(item.get("type") or "")
        if kind in {"function_call", "tool_call", "custom_tool_call"}:
            out.append(item)
            continue
        if kind == "message":
            for part in item.get("content") or []:
                if isinstance(part, dict) and part.get("type") in {"function_call", "tool_use", "tool_call"}:
                    name = part.get("name") or part.get("tool")
                    if name:
                        args = part.get("arguments") or part.get("input") or {}
                        out.append({"name": name, "arguments": args, "call_id": part.get("call_id") or part.get("id")})
    return out


_KV_LINE = re.compile(r'^[ \t]*([a-z][a-z0-9_]*)[ \t]*=[ \t]*"([^"]*)"[ \t]*$', re.I)
_TOOL_BLOCK = re.compile(
    r"(?m)^[ \t]*(?:invoke[ \t]+(?:tool[ \t]+)?)?([a-z][a-z0-9_]{1,48})[ \t]*\n"
    r"((?:[ \t]*[a-z][a-z0-9_]*[ \t]*=[ \t]*\"[^\"]*\"[ \t]*\n?)+)",
    re.I,
)
_XML_TAG = re.compile(
    r"<([a-z][a-z0-9_]{1,48})((?:\s+[a-z][a-z0-9_]*\s*=\s*\"[^\"]*\")+)\s*/>",
    re.I,
)
_CALL_FN = re.compile(
    r"(?m)^[ \t]*([a-z][a-z0-9_]{1,48})\(([^\n]{0,400})\)[ \t]*$",
    re.I,
)
_ATTR = re.compile(r'([a-z][a-z0-9_]*)\s*=\s*"([^"]*)"', re.I)


def _leaked_calls(text: str) -> list[dict[str, Any]]:
    """Parse tool dumps Grok sometimes writes as chat text instead of function_call items."""
    blob = (text or "").replace("\r\n", "\n").replace("\r", "\n").strip()
    if not blob:
        return []
    found: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()

    def add(name: str, args: dict[str, Any]) -> None:
        tool = (name or "").strip()
        if tool not in HANDLERS or tool == "offer_choices":
            return
        key = (tool, json.dumps(args, sort_keys=True, default=str))
        if key in seen:
            return
        seen.add(key)
        found.append({"name": tool, "arguments": args})

    for match in _TOOL_BLOCK.finditer(blob):
        args: dict[str, Any] = {}
        for line in match.group(2).splitlines():
            kv = _KV_LINE.match(line)
            if kv:
                args[kv.group(1)] = kv.group(2)
        add(match.group(1), args)
    for match in _XML_TAG.finditer(blob):
        args = {k: v for k, v in _ATTR.findall(match.group(2))}
        add(match.group(1), args)
    for match in _CALL_FN.finditer(blob):
        args = {k: v for k, v in _ATTR.findall(match.group(2))}
        if args:
            add(match.group(1), args)
    return found


def _strip_leaked(text: str) -> str:
    blob = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    blob = _TOOL_BLOCK.sub("", blob)
    blob = _XML_TAG.sub("", blob)

    def drop_call(match: re.Match[str]) -> str:
        name = match.group(1)
        args = {k: v for k, v in _ATTR.findall(match.group(2))}
        if name in HANDLERS and args:
            return ""
        return match.group(0)

    blob = _CALL_FN.sub(drop_call, blob)
    blob = re.sub(r"\n{3,}", "\n\n", blob)
    return blob.strip()


def _collect_calls(resp: dict[str, Any]) -> tuple[list[dict[str, Any]], bool]:
    native = _fn_calls(resp)
    if native:
        return native, False
    leaked = _leaked_calls(_text(resp))
    if not leaked:
        return [], False
    fake = []
    for i, item in enumerate(leaked):
        fake.append(
            {
                "type": "function_call",
                "name": item["name"],
                "arguments": json.dumps(item["arguments"]),
                "call_id": f"leak-{i}",
                "_leaked": True,
            }
        )
    return fake, True


def _text(resp: dict[str, Any]) -> str:
    bits: list[str] = []
    for item in resp.get("output") or []:
        if not isinstance(item, dict) or item.get("type") != "message":
            continue
        for part in item.get("content") or []:
            if isinstance(part, dict) and part.get("type") in ("output_text", "text") and part.get("text"):
                bits.append(str(part["text"]))
            elif isinstance(part, str):
                bits.append(part)
    joined = "\n".join(bits).strip()
    if joined:
        return joined
    ot = resp.get("output_text")
    if isinstance(ot, str) and ot.strip():
        return ot.strip()
    return ""


def _strip_dashes(text: str) -> str:
    import re

    text = text.replace("—", ". ").replace("–", ", ").replace(" - ", ", ")
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"`+", "", text)
    return text


def _grok_fail(err: str) -> str:
    low = err.lower()
    if "incorrect api key" in low or "invalid api key" in low or "401" in low:
        return "That Grok key is no good. Stick a real one in Settings and we are in business."
    if "429" in low or "rate" in low:
        return "Grok is swamped. Give it a second."
    return "Grok went quiet on me. Try again in a second."


_CRASH_QUIPS = (
    "Heh. The API just face-planted. Not me. Hit send again.",
    "I coughed up a hairball on the server. One more time.",
    "That wasn't a leak. That was a 500. Try me again.",
    "The rack hid under the desk. Ping me again.",
)


def chat_crash(exc: BaseException | str) -> str:
    text = str(exc)
    low = text.lower()
    if "internal server error" in low or "500" in low or not text.strip():
        import hashlib

        pick = int(hashlib.sha1(text.encode("utf-8", "replace")).hexdigest()[:8], 16) % len(_CRASH_QUIPS)
        return _CRASH_QUIPS[pick]
    return f"Heh. That one slipped. {text[:140]}"


def _run_tool(db: Session, user: User, name: str, args: dict[str, Any]) -> tuple[Any, bool]:
    import inspect

    fn = HANDLERS.get(name)
    if fn is None:
        return {"error": f"unknown tool {name}"}, False
    kwargs = dict(args)
    params = inspect.signature(fn).parameters
    if "db" in params:
        kwargs["db"] = db
    if "_db" in params:
        kwargs["_db"] = db
    if "user" in params:
        kwargs["user"] = user
    if "_user" in params:
        kwargs["_user"] = user
    extra = [key for key in list(kwargs) if key not in params and not any(p.kind is inspect.Parameter.VAR_KEYWORD for p in params.values())]
    for key in extra:
        kwargs.pop(key, None)
    result = fn(**kwargs)
    ok = not (isinstance(result, dict) and result.get("ok") is False)
    return result, ok


def _take_ui(result: Any, ui: list[dict[str, Any]]) -> None:
    if not isinstance(result, dict):
        return
    raw = result.get("ui")
    if isinstance(raw, dict):
        ui.append(raw)
    elif isinstance(raw, list):
        ui.extend(item for item in raw if isinstance(item, dict))


_CLUSTER_TABS = {"overview", "racks", "floor"}
_CLUSTER_SCENES = {"overview", "health", "racks", "floor", "cluster-name"}


def _coalesce_ui(cmds: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Power talk stays on Dynamic Power. Status tools must not yank back to cluster."""
    if not cmds:
        return cmds
    if not any(
        c.get("kind") == "navigate" and str(c.get("path") or "").split("?")[0] == "/maxlps" for c in cmds
    ):
        return cmds
    out: list[dict[str, Any]] = []
    seen_nav = False
    highlighted = False
    for c in cmds:
        kind = c.get("kind")
        if kind == "navigate":
            path = str(c.get("path") or "").split("?")[0]
            if path != "/maxlps" or seen_nav:
                continue
            seen_nav = True
            out.append({**c, "path": "/maxlps"})
            continue
        if kind == "click" and str(c.get("target") or "").lower() in _CLUSTER_TABS:
            continue
        if kind == "highlight":
            tgt = str(c.get("target") or "").lower()
            if tgt in _CLUSTER_SCENES:
                continue
            highlighted = True
        out.append(c)
    if not highlighted:
        out.append(
            {
                "kind": "highlight",
                "action": "look",
                "target": "limiter",
                "label": "dynamic power",
                "ids": [],
            }
        )
    return out


def _take_choices(result: Any) -> list[dict[str, Any]]:
    if not isinstance(result, dict):
        return []
    raw = result.get("choices")
    if not isinstance(raw, list):
        return []
    return [row for row in raw if isinstance(row, dict) and row.get("label") and row.get("tool")]


def run_chat(
    db: Session,
    user: User,
    messages: list[dict[str, str]],
    context: dict[str, str] | None = None,
    execute: dict[str, Any] | None = None,
    attachment_ids: list[str] | None = None,
    auth_header: str = "",
) -> dict[str, Any]:
    token = _AUTH_HEADER.set(auth_header or "")
    try:
        return _run_chat(
            db,
            user,
            messages,
            context=context,
            execute=execute,
            attachment_ids=attachment_ids,
        )
    finally:
        _AUTH_HEADER.reset(token)


def _run_chat(
    db: Session,
    user: User,
    messages: list[dict[str, str]],
    context: dict[str, str] | None = None,
    execute: dict[str, Any] | None = None,
    attachment_ids: list[str] | None = None,
) -> dict[str, Any]:
    key = current_key(db)
    if not key:
        return {
            "reply": "Hey! No Grok key yet. Stick one in Settings and we are in business.",
            "enabled": False,
            "actions": [],
            "choices": [],
            "ui": [],
        }
    history: list[dict[str, Any]] = []
    for m in messages[-12:]:
        role = m.get("role")
        content = (m.get("content") or "").strip()
        if role in ("user", "assistant") and content:
            history.append({"role": role, "content": content})

    atts = _user_attachments(db, user, attachment_ids) if attachment_ids else []
    image_parts: list[dict[str, Any]] = []
    text_bits: list[str] = []
    for row in atts:
        if row.kind == "image":
            part = _image_part(row)
            if part:
                image_parts.append(part)
            else:
                text_bits.append(f"Attached image {row.filename} could not be opened.")
        else:
            body = (row.text or "").strip()
            if body:
                text_bits.append(f"Attached file {row.filename}:\n{body[:12000]}")
            else:
                text_bits.append(f"Attached file {row.filename} ({row.kind}). No extractable text.")

    if not history or history[-1]["role"] != "user":
        if atts:
            history.append({"role": "user", "content": "Please read the attached file(s)."})
        else:
            return {"reply": "Say something first.", "enabled": True, "actions": [], "choices": [], "ui": []}

    last = history[-1]
    user_text = str(last.get("content") or "").strip()
    if text_bits:
        user_text = (user_text + "\n\n" if user_text else "") + "\n\n".join(text_bits)
    if image_parts:
        last["content"] = image_parts + [{"type": "input_text", "text": user_text or "Please read the attached image(s)."}]
    else:
        last["content"] = user_text or "Please read the attached file(s)."

    deep = _needs_reasoning(user_text, attachments=bool(atts), execute=execute)
    model = operator_model(user_text, attachments=bool(atts), execute=execute)
    smalltalk = _is_smalltalk(user_text) and not execute and not atts and not deep
    extra = ""
    path = (context or {}).get("path") if context else None
    if path:
        extra += f"\nThe user is currently looking at {path}."
    extra += _brief_cluster(db)
    sops = tool_list_sops(db).get("docs") or []
    if sops:
        extra += " SOP docs on file: " + ", ".join(str(d.get("title") or d.get("filename")) for d in sops) + "."
    if atts:
        extra += " Attached this turn: " + ", ".join(f"{r.filename} ({r.kind})" for r in atts) + "."

    actions: list[dict[str, Any]] = []
    mutated = False
    choices: list[dict[str, Any]] = []
    ui: list[dict[str, Any]] = []

    if execute:
        tool = str(execute.get("tool") or "").strip()
        label = str(execute.get("label") or tool).strip()
        raw_args = execute.get("arguments") if isinstance(execute.get("arguments"), dict) else {}
        if tool and tool != "offer_choices" and tool in HANDLERS:
            try:
                result, ok = _run_tool(db, user, tool, raw_args)
                if tool == "call_app":
                    mutated = mutated or str(raw_args.get("method") or "GET").upper() != "GET"
                elif tool in MUTATING:
                    mutated = True
            except PermissionError as exc:
                result, ok = {"error": str(exc)}, False
            except TypeError as exc:
                result, ok = {"error": str(exc)}, False
            except Exception as exc:
                result, ok = {"error": str(exc)}, False
            result = _attach_console_ui(tool, raw_args, result)
            actions.append({"name": tool, "ok": ok, "summary": _dump(result, 400)})
            _take_ui(result, ui)
            picked = _take_choices(result)
            if picked:
                choices = picked
            extra += (
                f"\nThe user tapped '{label}'. Already executed {tool}. "
                f"Result: {_dump(result, 2500)}. Report what happened, then offer_choices. "
                f"Do not call {tool} again."
            )
        elif tool:
            extra += f"\nThe user tapped '{label}' but {tool} is not a tool I can run."

    pasted = _pasted_add_node(user_text) if not execute else None
    if pasted:
        try:
            result, ok = _run_tool(db, user, "add_node", pasted)
            mutated = True
        except PermissionError as exc:
            result, ok = {"error": str(exc)}, False
        except Exception as exc:
            result, ok = {"error": str(exc)}, False
        actions.append({"name": "add_node", "ok": ok, "summary": _dump(_redact(result), 400)})
        _take_ui(result, ui)
        picked = _take_choices(result)
        if picked:
            choices = picked
        extra += (
            f"\nThe operator pasted a node. Already executed add_node. "
            f"Result: {_dump(_redact(result), 2500)}. Tell them it is on Provision. "
            f"Do not dump passwords. Do not call add_node again."
        )

    store = not bool(image_parts)
    payload: dict[str, Any] = {
        "model": model,
        "instructions": INSTRUCTIONS + extra,
        "input": history,
        "store": store,
        "max_output_tokens": 240 if smalltalk else (1800 if deep else 480),
    }
    if not smalltalk:
        payload["tools"] = TOOLS
        payload["tool_choice"] = "auto"
    try:
        resp = _xai_post(payload, key)
    except RuntimeError as exc:
        return {
            "reply": _grok_fail(str(exc)),
            "enabled": True,
            "actions": actions,
            "choices": _fallback_choices(db, actions),
            "mutated": mutated,
            "ui": _coalesce_ui(ui),
        }
    running_input: list[dict[str, Any]] = list(history)
    for _ in range(_MAX_ROUNDS):
        calls, leaked_round = _collect_calls(resp)
        if not calls:
            break
        outputs = []
        leak_notes: list[str] = []
        for call in calls:
            name = call.get("name") or ""
            raw_args = call.get("arguments") or "{}"
            try:
                args = json.loads(raw_args) if isinstance(raw_args, str) else dict(raw_args)
            except json.JSONDecodeError:
                args = {}
            if not isinstance(args, dict):
                args = {}
            try:
                result, ok = _run_tool(db, user, name, args)
                if name == "call_app":
                    mutated = mutated or str(args.get("method") or "GET").upper() != "GET"
                elif name in MUTATING:
                    mutated = True
            except PermissionError as exc:
                result = {"error": str(exc)}
                ok = False
            except TypeError as exc:
                result = {"error": str(exc)}
                ok = False
            except Exception as exc:
                result = {"error": str(exc)}
                ok = False
            result = _attach_console_ui(name, args, result)
            actions.append({"name": name, "ok": ok, "summary": _dump(result, 400)})
            _take_ui(result, ui)
            picked = _take_choices(result)
            if picked:
                choices = picked
            out_item = {
                "type": "function_call_output",
                "call_id": call.get("call_id"),
                "output": _dump(result),
            }
            outputs.append(out_item)
            if leaked_round:
                leak_notes.append(f"{name}: {_dump(result, 500)}")
            else:
                running_input.append(call)
                running_input.append(out_item)
        try:
            follow: dict[str, Any] = {
                "model": model,
                "tools": TOOLS,
                "max_output_tokens": 1800 if deep else 480,
            }
            if leaked_round:
                follow["instructions"] = (
                    INSTRUCTIONS
                    + extra
                    + "\nYou wrote tool calls as chat text. Already executed:\n"
                    + "\n".join(leak_notes)
                    + "\nSpeak to the user now in Mona voice. Do not write tool names or target=/label=/action= as text."
                )
                follow["input"] = history
                follow["store"] = False
            elif store and resp.get("id"):
                follow["previous_response_id"] = resp.get("id")
                follow["input"] = outputs
            else:
                follow["instructions"] = INSTRUCTIONS + extra
                follow["store"] = False
                follow["input"] = running_input
            resp = _xai_post(follow, key)
        except RuntimeError as exc:
            return {
                "reply": _grok_fail(str(exc)),
                "enabled": True,
                "actions": actions,
                "choices": choices or _fallback_choices(db, actions),
                "mutated": mutated,
                "ui": _coalesce_ui(ui),
            }
    reply = _strip_leaked(_strip_dashes(_text(resp) or ""))
    leftover = _leaked_calls(reply)
    if leftover:
        for item in leftover:
            try:
                result, ok = _run_tool(db, user, item["name"], item["arguments"])
            except Exception as exc:
                result, ok = {"error": str(exc)}, False
            result = _attach_console_ui(item["name"], item["arguments"], result)
            actions.append({"name": item["name"], "ok": ok, "summary": _dump(result, 400)})
            _take_ui(result, ui)
            picked = _take_choices(result)
            if picked:
                choices = picked
        reply = _strip_leaked(reply)
    if not reply.strip():
        reply = "Heh. On it."
    return {
        "reply": reply,
        "enabled": True,
        "actions": actions,
        "choices": choices or _fallback_choices(db, actions),
        "mutated": mutated,
        "model": model,
        "ui": _coalesce_ui(ui),
    }
